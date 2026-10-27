from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "IBM_AML_Interpretation_and_Review_Guide.xlsx"

NAVY = "17365D"
BLUE = "D9EAF7"
PALE_YELLOW = "FFF2CC"
PALE_GREEN = "E2F0D9"
LIGHT_GREY = "F2F2F2"
WHITE = "FFFFFF"
MID_GREY = "666666"
THIN_GREY = Side(style="thin", color="D9E1F2")


def add_title(ws, title, subtitle, last_column):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    cell = ws.cell(1, 1, title)
    cell.font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    cell = ws.cell(2, 1, subtitle)
    cell.font = Font(name="Aptos", size=10, italic=True, color=MID_GREY)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30


def add_section(ws, row, title, text, last_column, fill=PALE_YELLOW):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
    cell = ws.cell(row, 1, f"{title}  |  {text}")
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Aptos", size=10, bold=True, color=NAVY)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 36


def add_table_sheet(wb, name, title, subtitle, headers, rows, widths, table_name):
    ws = wb.create_sheet(name)
    add_title(ws, title, subtitle, len(headers))
    header_row = 4
    for column, label in enumerate(headers, start=1):
        cell = ws.cell(header_row, column, label)
        cell.font = Font(name="Aptos", bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN_GREY)
    ws.row_dimensions[header_row].height = 32

    for row_index, values in enumerate(rows, start=header_row + 1):
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row_index, column, value)
            cell.font = Font(name="Aptos", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN_GREY)
        ws.row_dimensions[row_index].height = 48

    last_row = header_row + len(rows)
    if rows:
        table = Table(displayName=table_name, ref=f"A{header_row}:{get_column_letter(len(headers))}{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = "1:4"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
    return ws


def build_workbook():
    wb = Workbook()
    readme = wb.active
    readme.title = "Start Here"
    readme.sheet_view.showGridLines = False
    readme.column_dimensions["A"].width = 27
    readme.column_dimensions["B"].width = 105
    add_title(
        readme,
        "IBM AML Data | Interpretation & Review Workbook",
        "A practical companion to IBM_AML_Transaction_Analysis_Procedure.docx. "
        "Use this workbook to interpret fields, follow the analysis procedure, "
        "and document review evidence.",
        2,
    )

    readme_rows = [
        ("What this is", "IBM synthetic transaction-monitoring data. It simulates transfers and includes an IBM simulator label; it is not a real bank ledger."),
        ("How to begin", "Read the Procedure tab, then use the Field Guide to interpret source fields and staged aliases. Use Rule Concepts only as review hypotheses."),
        ("How to follow a payment", "Start with Dataset Variant and Transaction Key, then record timestamp, sender bank/account, receiver bank/account, debit/credit, both currencies, format, and the label."),
        ("How to identify owners", "Match sender and receiver separately to the account data using Dataset Variant + normalized bank ID + account number. Check unmatched records."),
        ("How to use flags", "A rule flag means review the activity. It is not a finding of wrongdoing. Record supporting rows and context."),
        ("How to evaluate", "Compare alerts with Is Laundering only as a synthetic benchmark. For a rare positive label, report precision, recall, F1 and alert counts, not accuracy alone."),
    ]
    for row_index, (label, text) in enumerate(readme_rows, start=4):
        readme.cell(row_index, 1, label)
        readme.cell(row_index, 2, text)
        readme.cell(row_index, 1).font = Font(name="Aptos", bold=True, color=NAVY)
        readme.cell(row_index, 1).fill = PatternFill("solid", fgColor=BLUE)
        for cell in readme[row_index]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN_GREY)
        readme.row_dimensions[row_index].height = 43
    add_section(
        readme,
        11,
        "Critical safeguards",
        "Do not compare amounts across currencies without approved FX conversion. "
        "Do not join on account number alone. Preserve leading-zero identifiers. "
        "Do not treat simulator labels or rule alerts as proof.",
        2,
        PALE_YELLOW,
    )
    readme.freeze_panes = "A4"
    readme.sheet_properties.pageSetUpPr.fitToPage = True
    readme.page_setup.fitToWidth = 1
    readme.page_setup.fitToHeight = 1
    readme.page_setup.orientation = "landscape"

    fields = [
        ("Timestamp", "Source", "When the transaction was recorded.", "Sort chronologically and define explicit monitoring windows."),
        ("From Bank", "Source", "Bank/institution identifier for the sender.", "Code, not necessarily a bank name; preserve as text."),
        ("Account", "Source", "Sender account identifier.", "Combine with bank ID and variant when matching ownership."),
        ("To Bank", "Source", "Bank/institution identifier for the receiver.", "Code, not necessarily a bank name; preserve as text."),
        ("Account.1", "Source", "Receiver account identifier; repeated Account header in source.", "Use with To Bank for receiver ownership match."),
        ("Amount Paid", "Source", "Sender-side debit amount.", "Interpret in Payment Currency."),
        ("Payment Currency", "Source", "Currency of the amount paid by sender.", "Keep totals separate by currency unless using approved FX conversion."),
        ("Amount Received", "Source", "Receiver-side credit amount.", "Interpret in Receiving Currency."),
        ("Receiving Currency", "Source", "Currency of the amount received.", "May differ from Payment Currency."),
        ("Payment Format", "Source", "Recorded transfer format/channel.", "Contextual feature; a format by itself is not proof of risk."),
        ("Is Laundering", "Source", "IBM simulator transaction label: 1 or 0.", "Evaluation reference only; not a real-world determination."),
        ("Dataset Variant", "Staged", "Source sample identity (HI/LI and Small/Medium/Large).", "Keep in all joins, grouping, filters and outputs."),
        ("Transaction Key", "Staged", "Generated identifier based on variant and source row order.", "Traceability key; not an IBM transaction identifier."),
        ("Sender Account", "Staged", "Alias of Account.", "Convenient AML role name; source field remains available."),
        ("Receiver Account", "Staged", "Alias of Account.1.", "Convenient AML role name; source field remains available."),
        ("Debit Amount", "Staged", "Alias of Amount Paid.", "Sender outflow; same currency as Payment Currency."),
        ("Credit Amount", "Staged", "Alias of Amount Received.", "Receiver inflow; same currency as Receiving Currency."),
        ("Bank Name", "Account source", "Source-provided institution name in the account file.", "Use account dimension if display names are needed; not in transaction rows."),
        ("Bank ID", "Account source", "Bank identifier associated with the account.", "Normalize only for matching; retain the source text."),
        ("Account Number", "Account source", "Account identifier in the account dimension.", "Not unique by itself across banks."),
        ("Entity ID / Entity Name", "Account source", "Party/entity identifier and owner name.", "Use for entity-level grouping after validated joins."),
    ]
    add_table_sheet(
        wb,
        "Field Guide",
        "Field Guide",
        "Source columns and staged aliases translated into AML language.",
        ["Field", "Where", "Meaning", "Interpretation / control"],
        fields,
        [27, 18, 43, 66],
        "FieldGuide",
    )

    procedure = [
        (1, "Set scope", "Choose the dataset variant and date range. Record source/staged filenames, filters and analysis purpose."),
        (2, "Check quality", "Read manifest and quality JSON: row count, date range, nulls, label counts, expected columns and file paths."),
        (3, "Preserve identifiers", "Keep bank/account identifiers as text; do not discard leading zeroes. Confirm timestamps and amounts parse correctly."),
        (4, "Interpret movement", "Treat each row as sender-to-receiver; identify outgoing debit and incoming credit with separate currencies."),
        (5, "Link parties", "Join sender and receiver separately by variant + numeric-normalized bank ID + account number. Check match rates."),
        (6, "Summarize behavior", "Calculate account/entity inflows, outflows, volume, frequency, counterparties and timing within stated windows and currencies."),
        (7, "Define hypotheses", "Use documented patterns (e.g. burst, unusual amount, rapid onward flow, circular path) with calibrated thresholds."),
        (8, "Review evidence", "Trace all supporting transaction keys in time order; include owners, currencies, methods, alternatives and limitations."),
        (9, "Evaluate", "Compare to simulator label separately; report precision, recall, F1, false alerts and alert volume."),
        (10, "Explain outcome", "Write what was observed, why it merits review, what is uncertain, and the next review action. Do not claim guilt."),
    ]
    add_table_sheet(
        wb,
        "Procedure",
        "Analysis Procedure",
        "Follow in order; document exceptions rather than silently skipping checks.",
        ["Step", "Action", "What to do / record"],
        procedure,
        [10, 26, 105],
        "ProcedureSteps",
    )

    rules = [
        ("Unusual amount", "Amount outside the account's established range.", "Account + currency + period", "History, comparable transactions, currency, related movement.", "Use account-specific baseline; size alone is not proof."),
        ("Frequency burst", "Transaction count or value rises sharply in a short window.", "Account/entity + time window", "Chronological keys, baseline period, counterparties and spacing.", "Calibrate window and threshold; account for normal business cycles."),
        ("Many new counterparties", "Abrupt increase in distinct recipients/senders or banks.", "Account/entity + counterparties", "New vs known counterparties, linked entities, timing.", "Requires enough historical coverage; dataset start may truncate history."),
        ("Rapid pass-through", "Incoming value followed by outgoing movement soon after.", "Account + time-ordered flows", "Inbound and outbound records, time gap, currency and amounts.", "Amounts may not match exactly; avoid assuming funds are identical."),
        ("Circular movement", "A path appears to return to an earlier account/entity.", "Network path + time window", "Every hop, timestamps, accounts/entities, amounts/currencies.", "Validate graph/path logic and alternative unrelated transfers."),
        ("Repeated small transfers", "Frequent smaller transfers aggregate into a notable flow.", "Account/entity + currency + window", "All contributing transactions, sum, counterparties and period.", "Avoid arbitrary regulatory thresholds; none is supplied by this dataset."),
        ("Format anomaly", "Payment format departs from that account's observed behavior.", "Account + payment format", "Account history and relevant transaction context.", "Format alone is not a sufficient alert rationale."),
    ]
    add_table_sheet(
        wb,
        "Rule Concepts",
        "AML Rule Concepts",
        "These are explainable review hypotheses, not validated conclusions or universal thresholds.",
        ["Pattern", "Possible indicator", "Group by", "Evidence to inspect", "Caution"],
        rules,
        [25, 45, 34, 58, 55],
        "RuleConcepts",
    )

    case_fields = [
        ("Review reference", "Analyst-entered case or alert reference"),
        ("Dataset Variant", "Required; keep scope explicit"),
        ("Rule / hypothesis", "Name and version of rule that caused review"),
        ("Analysis window", "Start/end time and timezone, if known"),
        ("Sender bank / account", "From Bank + Account"),
        ("Sender entity", "Entity Name after validated join; note if unmatched"),
        ("Receiver bank / account", "To Bank + Account.1"),
        ("Receiver entity", "Entity Name after validated join; note if unmatched"),
        ("Transaction Keys", "All related keys, chronological order"),
        ("Amounts and currencies", "Debit + Payment Currency; Credit + Receiving Currency"),
        ("Payment format", "Value(s) and whether unusual for this account"),
        ("Observed behavior", "Describe facts, not conclusions"),
        ("Why review is warranted", "Explain the combination of indicators"),
        ("Alternative explanations", "Plausible benign explanations checked"),
        ("Data limitations", "Unmatched joins, missing history, FX or scope limitations"),
        ("IBM label", "Record separately for benchmark evaluation only"),
        ("Review disposition", "Project review status; no legal finding"),
        ("Next action", "Specific follow-up or reason no further action is proposed"),
    ]
    add_table_sheet(
        wb,
        "Case Review",
        "Investigation Review Template",
        "One row per case. Keep an additional transaction-detail tab/export when many transaction keys support one case.",
        ["Case field", "What to enter"],
        case_fields,
        [34, 100],
        "CaseReviewTemplate",
    )
    case_ws = wb["Case Review"]
    for row in range(5, 5 + len(case_fields)):
        case_ws.cell(row, 2).fill = PatternFill("solid", fgColor="FFFDF2")

    quality = [
        ("HI-Small transactions", "5,078,345", "Rows shown in HI-Small_Trans_quality.json.", "Use as a concrete example only; inspect report for chosen variant."),
        ("HI-Small Is Laundering = 1", "5,177", "Positive simulator-labeled rows in this quality report.", "Rare label; accuracy alone is misleading."),
        ("HI-Small Is Laundering = 0", "5,073,168", "Zero-labeled rows in this quality report.", "Class balance differs by dataset variant; do not assume all samples match."),
        ("HI-Small positive share", "About 0.10%", "5,177 divided by 5,078,345.", "Report precision, recall, F1 and alert workload."),
        ("Transaction key", "Generated", "Variant plus source row order; not IBM-issued.", "Useful for traceability inside this staged output."),
        ("Account join", "Composite", "Variant + normalized bank ID + account number.", "Do not join on account number alone."),
        ("Currency", "Separate per side", "Debit follows Payment Currency; credit follows Receiving Currency.", "No common FX conversion is defined here."),
        ("Bank identity", "Codes + dimension", "Transactions use bank IDs; account data includes Bank Name.", "Do not infer names from IDs."),
        ("Ground truth", "Synthetic label", "IBM simulator's transaction label.", "Not a real account/person determination."),
    ]
    add_table_sheet(
        wb,
        "Quality & Limits",
        "Quality Facts & Limitations",
        "The counts below come from the available HI-Small quality report; always check the report for the actual selected variant.",
        ["Topic", "Reported value", "What it means", "Analytical consequence"],
        quality,
        [29, 27, 62, 68],
        "QualityLimits",
    )
    qws = wb["Quality & Limits"]
    for row in range(5, 5 + len(quality)):
        if row in (6, 8, 9, 10, 11, 12, 13):
            for cell in qws[row]:
                cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)

    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddHeader.center.text = "&BIBM AML Interpretation Workbook"
        ws.oddFooter.center.text = "Synthetic-data analysis guide | Page &P of &N"
        ws.sheet_view.zoomScale = 85

    wb.properties.title = "IBM AML Interpretation and Review Guide"
    wb.properties.subject = "Spreadsheet companion explaining IBM AML fields, workflow, rules and investigation evidence"
    wb.properties.creator = "AML Analysis Project"
    wb.save(OUTPUT)

    check = load_workbook(OUTPUT, read_only=False, data_only=False)
    expected = {"Start Here", "Field Guide", "Procedure", "Rule Concepts", "Case Review", "Quality & Limits"}
    if set(check.sheetnames) != expected:
        raise RuntimeError(f"Unexpected workbook sheets: {check.sheetnames}")
    if check["Field Guide"].max_row < 20 or check["Procedure"].max_row != 14:
        raise RuntimeError("Workbook validation failed: required guide content is missing.")
    print(f"Created and validated {OUTPUT}")
    print(f"Sheets: {', '.join(check.sheetnames)}")


if __name__ == "__main__":
    build_workbook()
