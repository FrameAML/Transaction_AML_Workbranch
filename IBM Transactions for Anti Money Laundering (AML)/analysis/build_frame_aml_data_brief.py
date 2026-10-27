from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "FRAME_IBM_AML_Data_Use_Brief.docx"

NAVY = "17365D"
BLUE = "D9EAF7"
PALE_YELLOW = "FFF2CC"
PALE_GREEN = "E2F0D9"
WHITE = "FFFFFF"
GREY = "595959"


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell(cell, text: str, *, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(8.3)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_table(
    document: Document,
    headers: list[str],
    rows: list[tuple[str, ...]],
    widths: list[float],
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_props = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_props.append(repeat)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell(cell, header, bold=True, color=WHITE)
        shade_cell(cell, NAVY)
        cell.width = Inches(widths[index])
    for row_data in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row_data):
            set_cell(cells[index], text)
            cells[index].width = Inches(widths[index])
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.add_run(text)


def add_callout(document: Document, heading: str, text: str, fill: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(f"{heading}: ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    paragraph.add_run(text)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("FRAME | IBM AML synthetic-data brief  •  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GREY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build_document() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = RGBColor.from_string("202124")
    styles["Normal"].paragraph_format.space_after = Pt(3)
    styles["Normal"].paragraph_format.line_spacing = 1.0
    for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 10)):
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(3)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("FRAME CONSULTING  |  ANALYTICS CONCEPT")
    header_run.font.name = "Aptos"
    header_run.font.size = Pt(8)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor.from_string(NAVY)
    add_page_number(section.footer.paragraphs[0])

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(0)
    title.add_run("IBM AML Data: Use & Pattern Brief")
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(7)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run(
        "A practical proposal for Frame: what the synthetic dataset can support, "
        "how to structure detection patterns, and what not to infer."
    )
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(GREY)

    add_callout(
        document,
        "Executive point",
        "Use this as a sandbox for demonstrating transaction monitoring, "
        "prototyping explainable rules and graph analytics, and testing evaluation "
        "workflows. It is synthetic research data—not evidence about real customers "
        "and not, by itself, a validated production AML control.",
        PALE_GREEN,
    )

    document.add_heading("1. What this dataset represents", level=1)
    document.add_paragraph(
        "IBM describes a synthetic financial ecosystem of people, businesses, "
        "banks, accounts and transactions. The generator simulates placement, "
        "layering and integration, tracks illicit funds across multiple transfers, "
        "and supplies transaction-level labels. This makes the data useful for "
        "controlled experiments where real transaction data and complete ground "
        "truth are difficult to obtain [1, 2]."
    )
    add_bullet(
        document,
        "Six independent variants: HI-Small/Medium/Large and LI-Small/Medium/Large. "
        "HI has a higher illicit ratio than LI; Small/Medium/Large are different "
        "scales, not nested samples [1].",
    )
    add_bullet(
        document,
        "A transaction is a directed movement: sender bank/account → receiver "
        "bank/account, with timestamp, paid and received amounts, their separate "
        "currencies, payment format, and Is Laundering label.",
    )
    add_bullet(
        document,
        "Account files add bank and entity-owner attributes. Join each side with "
        "Dataset Variant + normalized Bank ID + Account Number; account number "
        "alone is not a safe key.",
    )
    add_bullet(
        document,
        "Pattern files document eight generator scenarios: FAN-IN, FAN-OUT, CYCLE, "
        "BIPARTITE, GATHER-SCATTER, SCATTER-GATHER, STACK and RANDOM. These are "
        "useful scenario references, not an exhaustive list of laundering behavior [1].",
    )

    document.add_heading("Dataset scale in the local staged copy", level=2)
    add_table(
        document,
        ["Variant", "Transactions", "IBM label = 1", "Positive rate"],
        [
            ("HI-Small", "5,078,345", "5,177", "0.1019%"),
            ("HI-Medium", "31,898,238", "35,230", "0.1104%"),
            ("HI-Large", "179,702,229", "225,546", "0.1255%"),
            ("LI-Small", "6,924,049", "3,565", "0.0515%"),
            ("LI-Medium", "31,251,483", "16,041", "0.0513%"),
            ("LI-Large", "176,066,557", "100,604", "0.0571%"),
        ],
        [1.15, 1.55, 1.35, 1.1],
    )
    document.add_paragraph(
        "Local totals: 430,920,901 transactions, 386,163 positive labels "
        "(about 0.09%), and 9,566,360 account rows. The extreme class imbalance "
        "means a model can report high accuracy while missing most positive cases."
    )
    document.add_heading("Account type is not supplied by IBM", level=2)
    document.add_paragraph(
        "The account file provides Bank Name, Bank ID, Account Number, Entity ID "
        "and Entity Name, but no account product/type, GL purpose, settlement role "
        "or balance. A clearing account is a temporary ledger account used to "
        "collect, route or reconcile postings; a settlement account discharges "
        "payment obligations between participants. Correspondent/Nostro/Vostro, "
        "suspense, transit, escrow/trust, merchant-settlement and pooled/omnibus "
        "roles likewise need authoritative account-master, GL, legal or payment-"
        "system evidence. High volume, multiple counterparties or rapid pass-through "
        "may prompt review but do not identify an account type. Product type, "
        "owner type and operational role can overlap; use separate, effective-dated "
        "fields and leave unsupported values unknown."
    )
    document.add_paragraph(
        "Use IBM's supplied Bank Name after joining each transaction endpoint to "
        "the same-variant account dimension with bank ID and account number. IBM "
        "Bank IDs are synthetic benchmark-local keys, not BIC/SWIFT codes or "
        "domestic routing numbers. For real-world enrichment, keep the identifier "
        "scheme and jurisdiction explicit and use verified, effective-dated "
        "directory records. See docs/ACCOUNT_AND_BANK_IDENTIFIER_GUIDE.md."
    )

    document.add_heading("2. Useful Frame applications", level=1)
    add_table(
        document,
        ["Workstream", "Practical output for Frame"],
        [
            (
                "Explainable proof of concept",
                "A reproducible demo that traces one transfer and shows why an "
                "account or transaction entered review.",
            ),
            (
                "Pattern library",
                "A reusable catalogue of rules, graph signals, evidence requirements, "
                "exceptions and review actions.",
            ),
            (
                "Analytics benchmarking",
                "Compare transparent rules, tabular baselines and graph approaches "
                "on synthetic labels; document precision/recall trade-offs.",
            ),
            (
                "Data and scale engineering",
                "Test joins, lineage, temporal features, alert evidence and processing "
                "on Small/Medium/Large files before designing client-scale architecture.",
            ),
        ],
        [1.6, 5.1],
    )
    add_callout(
        document,
        "Positioning",
        "Describe the result as a synthetic-data prototype or capability "
        "demonstration. Do not claim it detects real financial crime, validates "
        "a client’s controls, or establishes that a real person or transaction is suspicious.",
        PALE_YELLOW,
    )

    document.add_page_break()
    document.add_heading("3. Structure Frame’s pattern library", level=1)
    document.add_paragraph(
        "Represent each transaction as a time-stamped directed edge between account "
        "nodes. Group first by dataset variant, bank and account; add entity-owner "
        "views only after validated joins. The labels below are IBM scenario names "
        "with plain-language interpretations for analysis—not claims about formal "
        "regulatory typologies."
    )
    add_table(
        document,
        ["Scenario family", "Graph shape to test", "Candidate review signal"],
        [
            (
                "FAN-OUT",
                "One source sends to many destinations.",
                "Unusually high distinct receivers or outbound value in a short, "
                "calibrated window.",
            ),
            (
                "FAN-IN",
                "Many sources send to one destination.",
                "Unusually high distinct senders, inbound count or value to an account.",
            ),
            (
                "CYCLE",
                "Funds move along a time-respecting path back to an earlier node.",
                "Closed paths within bounded hops/time; preserve every transfer in "
                "the evidence trail.",
            ),
            (
                "BIPARTITE",
                "Transfers connect two sets of accounts.",
                "Dense or repeated cross-group flows, especially with rapid onward movement.",
            ),
            (
                "GATHER-SCATTER",
                "Funds converge, then disperse.",
                "Inbound aggregation followed by multiple outbound counterparties.",
            ),
            (
                "SCATTER-GATHER",
                "Funds disperse, then converge.",
                "A set of sources fans out and later routes value to a common receiver.",
            ),
            (
                "STACK",
                "Funds pass through a chain of accounts.",
                "Multiple short-interval hops, new counterparties or rapid pass-through.",
            ),
            (
                "RANDOM / other",
                "Generator-defined route without one of the named shapes.",
                "Keep as an annotated scenario class; do not invent a universal "
                "“random pattern” rule.",
            ),
        ],
        [1.35, 2.05, 3.3],
    )

    document.add_heading("Specify every Frame rule consistently", level=2)
    add_table(
        document,
        ["Rule field", "What to record"],
        [
            ("Rule ID and hypothesis", "Plain-language behavior being tested; distinguish signal from conclusion."),
            ("Scope and grouping", "Account/entity/network; variant; sender/receiver role; currency."),
            ("Window and conditions", "Lookback, count/value/degree/path logic, calibrated threshold and version."),
            ("Evidence", "Triggering transaction keys, timestamps, parties, linked hops and feature values."),
            ("Context and exceptions", "Expected baseline, known benign explanations, data gaps and join quality."),
            ("Action and evaluation", "Analyst question/action, reviewer disposition, precision/recall and alert workload."),
        ],
        [1.55, 5.15],
    )
    add_callout(
        document,
        "Threshold principle",
        "Start with account- and currency-specific baselines (for example, rolling "
        "count, distinct counterparties and amount deviation). Treat 1-hour/24-hour/"
        "7-day windows as candidate experiments, not fixed policy thresholds. "
        "Calibrate using evidence, review capacity and appropriate business context.",
        BLUE,
    )

    document.add_page_break()
    document.add_heading("4. A defensible detection and review workflow", level=1)
    add_table(
        document,
        ["Step", "Frame workflow"],
        [
            (
                "1. Validate",
                "Check schema, duplicates, missing values, timestamps, identifiers, "
                "label rates and account-join coverage. Preserve source row lineage.",
            ),
            (
                "2. Build behavior",
                "Create time-ordered account/entity activity: counts, value by "
                "currency, distinct counterparties, first-seen relationships, "
                "in/out-degree and short path/cycle features.",
            ),
            (
                "3. Generate signals",
                "Run versioned rules and, separately, baseline models. Keep IBM "
                "labels and pattern-file annotations out of model input features.",
            ),
            (
                "4. Explain alerts",
                "Show the triggering row, related transactions in time order, party "
                "and account identifiers, currencies, rule logic and limitations.",
            ),
            (
                "5. Test out of time",
                "Use chronological train/validation/test splits (IBM reports 60/20/20), "
                "and prevent future transactions from entering historical features [2].",
            ),
            (
                "6. Measure and refine",
                "Report minority-class precision, recall, F1/PR-AUC, false alerts "
                "and review workload. Tune thresholds to the intended use, not just "
                "to maximize a single score.",
            ),
        ],
        [1.25, 5.45],
    )

    document.add_heading("Critical interpretation and data caveats", level=1)
    add_bullet(
        document,
        "Kaggle says the published date ranges are primary periods and describes "
        "later transactions as laundering. However, this full local scan finds "
        "56,691 rows after the stated cutoffs across the six variants, of which "
        "33,583—not all—have Is Laundering = 1. For example, HI-Small has 1,108 "
        "post-cutoff rows and 655 positives. This does not reproduce the page's "
        "all-positive statement under the boundary used here. Reconcile the exact "
        "cutoff, data version and cited discussion before using a temporal split; "
        "never treat date alone as ground truth [1].",
    )
    add_bullet(
        document,
        "Do not use random row splits as the only test. Report primary-period and "
        "post-period results separately and retain the observed discrepancy as an "
        "open data-validation issue until reconciled.",
    )
    add_bullet(
        document,
        "Do not treat Is Laundering as real-world confirmed suspicion. It is "
        "synthetic benchmark ground truth. IBM also notes that not every labelled "
        "laundering transaction follows one of the eight named scenarios [1].",
    )
    add_bullet(
        document,
        "Do not compare or sum paid and received amounts across currencies without "
        "a documented FX rate, effective time and conversion method. IDs are "
        "synthetic; the generated Transaction Key in this project is a row-trace "
        "key, not an IBM-issued ID.",
    )
    add_bullet(
        document,
        "Keep account IDs out of predictive features when they can let a model "
        "memorize entities; prefer behavior and graph features. The paper’s "
        "experiments also exclude source/destination IDs as direct features [2].",
    )

    document.add_heading("5. What is actually exported here?", level=1)
    add_callout(
        document,
        "Short answer",
        "The complete six-variant source set is present locally and was staged: "
        "18 source files (transactions, accounts and pattern descriptions for each "
        "variant), 430,920,901 transaction rows and 9,566,360 account rows. The "
        "staging folder contains Parquet transaction/account tables, copied pattern "
        "text files, quality summaries and a manifest (about 17.32 GiB).",
        PALE_GREEN,
    )
    document.add_paragraph(
        "But IBM_AML_Analysis-updated.pbix is not the full data export. It is a "
        "Power BI report/model, separate from the full six-variant staging outputs. "
        "The project notes explicitly say to connect Power BI to selected Parquet "
        "tables; loading every Large variant into Desktop may exceed available "
        "memory. Use the Parquet outputs for full-data work and the PBIX for its "
        "report/model content. A generated Transaction Key is only a local row key."
    )

    document.add_heading("6. Related but separate: credit-card fraud data", level=1)
    document.add_paragraph(
        "The related Kaggle Credit Card Transactions page describes a separate IBM "
        "synthetic fraud dataset: more than 20 million purchase transactions from "
        "2,000 simulated US consumers, covering decades and multiple cards per "
        "consumer. Its current file list includes credit_card_transactions-ibm_v2.csv, "
        "User0_credit_card_transactions.csv, sd254_cards.csv and sd254_users.csv. "
        "This can support a separate card-fraud or purchase-pattern test suite, but "
        "fraud labels and card-purchase behavior are not AML laundering labels or "
        "inter-account transfer typologies. Do not combine its metrics with the IBM "
        "AML benchmark. The page description includes Apache License 2.0 terms while "
        "Kaggle metadata says 'Other (specified in description)'; verify terms "
        "before redistribution [5]."
    )

    document.add_heading("Sources and use notes", level=1)
    sources = [
        "[1] Kaggle, Erik Altman, “IBM Transactions for Anti Money Laundering "
        "(AML)” — dataset description, variants, class ratios, scenario files, "
        "date-range caveat and license. "
        "https://www.kaggle.com/datasets/ealtman2019/ibm-transactions-for-anti-money-laundering-aml",
        "[2] Altman et al., “Realistic Synthetic Financial Transactions for "
        "Anti-Money Laundering Models,” NeurIPS 2023 Datasets and Benchmarks. "
        "https://arxiv.org/abs/2306.16424",
        "[3] IBM, Multi-GNN AML models and experiment code. "
        "https://github.com/IBM/Multi-GNN",
        "[4] Dataset license shown by Kaggle: Community Data License Agreement "
        "(CDLA) – Sharing, Version 1.0. Review attribution and redistribution "
        "terms before republishing the data. https://cdla.dev/sharing-1-0/",
        "[5] Kaggle, IBM Credit Card Transactions — fraud benchmark description "
        "and current file inventory. "
        "https://www.kaggle.com/datasets/ealtman2019/credit-card-transactions",
        "[6] Kaggle community notebook, “Baseline - IBM Transactions for AML,” "
        "by Gaurav Bhattacharya. Public page metadata was checked, but notebook "
        "code and outputs were not retrievable in this review; no implementation "
        "or performance claims are attributed to it. "
        "https://www.kaggle.com/code/gbiamgaurav/baseline-ibm-transactions-for-aml",
    ]
    for source in sources:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(source)
        run.font.size = Pt(7.5)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(3)
    note.paragraph_format.space_after = Pt(0)
    run = note.add_run(
        "Prepared 26 October 2026 from the Kaggle dataset pages, linked research, "
        "and local staging manifest/quality reports. This is an analytical "
        "discussion brief, not legal advice or an operational AML policy."
    )
    run.italic = True
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string(GREY)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
