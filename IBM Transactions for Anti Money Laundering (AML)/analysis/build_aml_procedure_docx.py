from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "IBM_AML_Transaction_Analysis_Procedure.docx"

NAVY = "17365D"
BLUE = "D9EAF7"
PALE_YELLOW = "FFF2CC"
PALE_GREEN = "E2F0D9"
LIGHT_GREY = "F2F2F2"
WHITE = "FFFFFF"


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, *, bold=False, color=None, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_text(cell, header, bold=True, color=WHITE, size=9)
        shade_cell(cell, NAVY)
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            set_cell_text(cells[index], value)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_numbered(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(item)


def add_callout(document, title, text, fill=PALE_YELLOW):
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    heading = paragraph.add_run(f"{title}: ")
    heading.bold = True
    heading.font.color.rgb = RGBColor.from_string(NAVY)
    paragraph.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_step(document, title, purpose, bullets):
    document.add_heading(title, level=2)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    label = paragraph.add_run("Purpose: ")
    label.bold = True
    paragraph.add_run(purpose)
    add_bullets(document, bullets)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 28), ("Heading 1", 17), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.font.bold = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("IBM AML DATA  |  ANALYSIS PROCEDURE")
    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Working guide for analysis and training  |  Synthetic data; not a real-world case")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("666666")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(40)
    title.paragraph_format.space_after = Pt(5)
    run = title.add_run("IBM AML Transaction\nAnalysis Procedure")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("A practical, traceable guide for Francois and the project team")
    run.italic = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string("595959")

    add_callout(
        doc,
        "Purpose",
        "Follow a consistent procedure to understand the source data, connect transactions to account owners, identify review-worthy patterns, and explain findings in plain AML language.",
        PALE_GREEN,
    )
    add_callout(
        doc,
        "Important limitation",
        "This is IBM synthetic AML data. `Is Laundering` is the simulator's label, not proof of a crime or a real-world bank decision. An analytical flag is a lead for review, not a conclusion.",
    )

    doc.add_heading("Quick procedure", level=1)
    add_numbered(
        doc,
        [
            "Choose and record one dataset variant; do not mix variants without keeping the variant identifier.",
            "Check the staged manifest and quality report, then confirm the columns and date/amount fields.",
            "Interpret each transaction as a sender-to-receiver movement; preserve both amount/currency sides.",
            "Attach sender and receiver owner details using the full composite account key.",
            "Build behavior summaries within relevant time windows, currencies, and account/entity groups.",
            "Use pattern indicators to create explainable review alerts; calibrate thresholds rather than guessing them.",
            "Trace every alert back to its transaction keys and surrounding transactions; document evidence and limitations.",
            "Compare detection results with `Is Laundering` only as a simulator benchmark, then explain findings and next steps.",
        ],
    )

    doc.add_heading("1. Understand the files and select scope", level=1)
    add_bullets(
        doc,
        [
            "`archive/` contains the original IBM source data; keep it unchanged.",
            "`powerbi_staging/manifest.json` describes which variants and files were staged. Use the paths recorded there rather than assuming a particular subfolder.",
            "Transaction Parquet files contain the transaction fact rows; account Parquet files contain account/entity details; `patterns/` contains IBM scenario descriptions.",
            "The variants include HI/LI and Small/Medium/Large. Record the chosen variant and its size in every analysis output.",
            "For an initial learning/review exercise, a Small variant is easier to handle; use the larger variants only when resources and the analysis question justify it.",
        ],
    )
    add_callout(
        doc,
        "Trace key",
        "`Transaction Key` is a generated row identifier based on variant and source row order. It is useful for reproducing and finding a row, but it is not an IBM-issued transaction ID.",
        BLUE,
    )

    doc.add_heading("2. Read the transaction fields correctly", level=1)
    add_table(
        doc,
        ["Field(s)", "AML interpretation", "How to use"],
        [
            ("Timestamp", "When the movement was recorded.", "Sort into a consistent timeline and define any analysis window."),
            ("From Bank + Account", "Sender institution code and sender account.", "Identify the origin account; retain IDs as text."),
            ("To Bank + Account.1", "Receiver institution code and receiver account.", "Identify the destination account; retain IDs as text."),
            ("Amount Paid / Debit Amount", "Sender-side amount leaving the account.", "Analyze outgoing flow in its Payment Currency."),
            ("Amount Received / Credit Amount", "Receiver-side amount arriving in the account.", "Analyze incoming flow in its Receiving Currency."),
            ("Payment Currency / Receiving Currency", "Currency on each side of the transfer.", "Never add or compare amounts across currencies without a documented conversion."),
            ("Payment Format", "Recorded transfer format/channel.", "Treat as context; a format alone does not establish suspicious activity."),
            ("Is Laundering", "IBM simulator label: 1 or 0.", "Use as a reference label for evaluation, not as proof or as an operational decision."),
            ("Dataset Variant", "Identifies the source variant for a row.", "Keep in filters, keys, joins, grouping, and outputs."),
        ],
        [1.65, 2.45, 3.2],
    )
    add_callout(
        doc,
        "Do not assume",
        "Transaction bank fields are codes, not necessarily names. The account source includes a `Bank Name` field; use that dimension if a name is needed. Do not infer bank names from numeric IDs.",
    )

    doc.add_heading("3. Validate data before analysis", level=1)
    add_numbered(
        doc,
        [
            "Open the variant's transaction and account quality JSON reports and note row counts, source file names, date range, nulls, and label counts.",
            "Confirm expected transaction fields exist and that identifiers are still text. Leading zeroes are meaningful for display and must not be lost.",
            "Check that timestamps parse, amount fields are numeric, currency and payment-format values are present, and `Is Laundering` contains only expected 0/1 values.",
            "Confirm transaction keys are populated and unique within the staged variant.",
            "Record any missing data, duplicates, unexpected types, or filtering performed. Do not silently discard rows.",
        ],
    )
    add_callout(
        doc,
        "Imbalance warning",
        "HI-Small has 5,177 rows labelled 1 out of 5,078,345 (about 0.10%) in the available quality report. With a rare positive label, overall accuracy can look high even when a detector misses most positives. Review precision, recall, F1, and false-alert volume.",
    )

    doc.add_heading("4. Link each transaction to the right account owner", level=1)
    doc.add_paragraph(
        "Use separate sender-side and receiver-side matches to the account dimension. The account number by itself is not a safe key because it may repeat at different banks."
    )
    add_table(
        doc,
        ["Transaction side", "Transaction fields", "Account dimension fields"],
        [
            ("Sender", "`Dataset Variant`, `From Bank`, `Account`", "`Dataset Variant`, normalized `Bank ID`, `Account Number`"),
            ("Receiver", "`Dataset Variant`, `To Bank`, `Account.1`", "`Dataset Variant`, normalized `Bank ID`, `Account Number`"),
        ],
        [1.1, 3.05, 3.15],
    )
    add_bullets(
        doc,
        [
            "Normalize bank IDs only for matching (for example, `010` and `10` compare as the same numeric ID); preserve the original source value for display and audit.",
            "Keep `Dataset Variant` in the join key so records from different samples are not accidentally joined.",
            "After joining, check match rates and unmatched rows on both sides. Investigate unexpected unmatched records; do not silently treat them as having no owner.",
            "Use `Entity Name` for the account owner and `Bank Name` for the source-provided bank name when available.",
            "If a transaction table has two account joins, use clearly named sender-owner and receiver-owner fields to prevent confusion.",
        ],
    )

    doc.add_heading("5. Create interpretable behavior summaries", level=1)
    doc.add_paragraph(
        "Start with descriptive summaries before introducing a model. Group the data by account/entity and time, and preserve the currency context."
    )
    add_bullets(
        doc,
        [
            "Incoming and outgoing transaction counts and totals by account/entity and time window.",
            "Typical transaction amount and variability for each account, calculated separately by currency.",
            "Number of distinct counterparties and banks, including newly observed counterparties where history permits.",
            "Time between an incoming credit and subsequent outgoing debit, as a possible rapid-throughput indicator.",
            "Repeated sender-receiver relationships, high-frequency bursts, and paths where funds appear to return to an earlier account.",
            "Payment-format and currency-pair distributions for each account/entity.",
        ],
    )
    add_callout(
        doc,
        "Currency control",
        "The paid amount and received amount can be denominated in different currencies. Keep totals separated by currency unless an approved FX rate source, conversion timestamp, and conversion method are specified.",
        PALE_YELLOW,
    )

    doc.add_heading("6. Define review rules and alert logic", level=1)
    doc.add_paragraph(
        "Treat each item below as an investigation hypothesis. Set thresholds from observed baselines and operational capacity; do not present example concepts as proven laundering."
    )
    add_table(
        doc,
        ["Pattern hypothesis", "Possible indicator", "Evidence to inspect"],
        [
            ("Unusual amount", "Amount far outside this account's normal range.", "Account history, currency, comparable activity, related transfers."),
            ("Activity burst", "Transaction count/value rises sharply within a short window.", "Exact transactions, prior baseline, counterparties, time spacing."),
            ("Many counterparties", "Sudden increase in distinct recipients or sending parties.", "New vs established relationships, shared entities, timing."),
            ("Rapid pass-through", "Funds received and then sent onward soon after.", "Incoming/outgoing amount and currency, timing, linked transfer chain."),
            ("Circular movement", "Funds appear to return through a sequence of accounts.", "Full path, time order, amounts/currency; account for unrelated flows."),
            ("Repeated small transfers", "Many individually small transfers aggregate to notable flow.", "Aggregate by entity/account, time window, counterparty, currency."),
            ("Format anomaly", "Format differs from the account's own historical pattern.", "Customer/account baseline and transaction context; format alone is insufficient."),
        ],
        [1.35, 2.55, 3.45],
    )
    add_bullets(
        doc,
        [
            "Use a documented baseline and window for each rule; state whether thresholds are illustrative or calibrated.",
            "Avoid one universal monetary threshold across currencies, variants, or account types.",
            "Record which transactions triggered each rule and the feature values that caused the trigger.",
            "Combine indicators carefully and retain rule-level reasons; a composite score must remain explainable.",
            "Review a sample of non-alerted as well as alerted records when evaluating detection quality.",
        ],
    )

    doc.add_heading("7. Investigate and document each alert", level=1)
    add_numbered(
        doc,
        [
            "Capture the alert/rule name, run date, dataset variant, window, and rule version.",
            "List every supporting `Transaction Key` in chronological order, not just the first triggering row.",
            "Show sender and receiver bank/account codes plus linked entity names, where the account match succeeded.",
            "Show timestamp, debit and credit amounts, their respective currencies, payment format, and IBM label.",
            "Explain in plain language what pattern was observed and why it deserves review.",
            "State alternative benign explanations, data limitations, missing joins, and any currency treatment.",
            "Record the analyst disposition as a project-review outcome (for example, needs more context / not prioritized / escalated for further review). Do not claim a legal finding.",
        ],
    )
    add_callout(
        doc,
        "Minimum reproducibility",
        "A second analyst should be able to reproduce the alert from the variant, rule/version, time window, keys, and documented calculation. If not, the alert explanation is incomplete.",
        PALE_GREEN,
    )

    doc.add_heading("8. Evaluate rules using the IBM label", level=1)
    add_bullets(
        doc,
        [
            "Compare predicted alerts with `Is Laundering` only after the rule/features are defined; clearly label this as evaluation against simulator labels.",
            "Report true positives, false positives, false negatives, precision, recall, F1, and the number of alerts requiring review.",
            "Use a time-aware or scenario-aware split if training/tuning a model; avoid leaking the label into model features.",
            "Inspect performance across variants and relevant subgroups; a good aggregate score may hide poor coverage.",
            "Do not use a label-derived field as a detection input. Keep label-based evaluation separate from operational rule logic.",
        ],
    )
    add_callout(
        doc,
        "Interpretation",
        "The IBM label provides a benchmark for this synthetic dataset. It is not evidence that a real account or person committed money laundering.",
    )

    doc.add_heading("9. Final reporting format", level=1)
    add_table(
        doc,
        ["Report item", "What to state"],
        [
            ("Scope", "Variant, source/staged file, date range, row count, filters."),
            ("Method", "Join keys, rules, windows, baselines, currency handling, and rule version."),
            ("Finding", "What behavior was observed, counts/amounts by currency, and involved entities/accounts."),
            ("Trace", "Transaction keys and chronological transaction details."),
            ("Evaluation", "Metrics against IBM label and alert volume; explain limitations."),
            ("Conclusion", "What can be said from this analysis, what cannot, and recommended next review."),
        ],
        [1.45, 5.9],
    )

    doc.add_heading("Important facts at a glance", level=1)
    add_bullets(
        doc,
        [
            "One transaction row is a money movement from an origin account to a destination account.",
            "`Amount Paid`/`Debit Amount` is sender-side outflow; `Amount Received`/`Credit Amount` is receiver-side inflow.",
            "Always read amounts with their corresponding currency.",
            "Bank IDs in transactions are codes; use account data for source-provided bank names and entity ownership.",
            "Join by variant + normalized bank ID + account number, not account number alone.",
            "`Transaction Key` is generated for traceability; it is not a source-system identifier.",
            "`Is Laundering` is a simulator label, not a real-world determination.",
            "Unusual behavior creates a review lead; context, evidence, and human review are required.",
        ],
    )

    doc.add_heading("Project references", level=1)
    add_bullets(
        doc,
        [
            "`README.md` — project overview and folder map.",
            "`docs/AML_DATA_INTERPRETATION.md` — detailed field and join interpretation.",
            "`docs/FRANCOIS_AML_ANALYSIS_CHECKLIST.md` — concise case-review checklist.",
            "`PowerBI_Staging_IBM_AML_Dataset.md` — staging pipeline, source inventory, and join guidance.",
            "`stage_aml_for_powerbi.py` — reproducible source-to-Parquet staging process.",
            "`powerbi_staging/manifest.json` and the per-variant quality JSON — actual staged paths and quality facts.",
        ],
    )

    doc.core_properties.title = "IBM AML Transaction Analysis Procedure"
    doc.core_properties.subject = "Step-by-step interpretation and review procedure for IBM synthetic AML data"
    doc.core_properties.author = "AML Analysis Project"
    doc.core_properties.keywords = "AML, transaction monitoring, IBM synthetic data, interpretation"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_document()
