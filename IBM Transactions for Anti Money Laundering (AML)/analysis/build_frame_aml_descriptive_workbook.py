from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import re
from collections import Counter, defaultdict
from copy import copy
from datetime import date
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import numpy as np
import pandas as pd
import pyarrow.parquet as pq
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.table import TableColumn
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
STAGING = ROOT / "powerbi_staging"
ARCHIVE = ROOT / "archive"
OUTPUT_DIR = ROOT / "docs"
TEST_DIR = ROOT / "test_data"
MANIFEST_PATH = STAGING / "manifest.json"
PROFILE_PATH = ROOT / "analysis" / "frame_aml_full_population_profile.json"
WORKBOOK_PATH = OUTPUT_DIR / "FrameAML_IBM_Data_Preparation_and_Descriptive_Analysis.xlsx"
SUMMARY_WORKBOOK_PATH = OUTPUT_DIR / (
    "FrameAML_IBM_Data_Preparation_and_Descriptive_Analysis_Summary.xlsx"
)
RYNO_DIR = ROOT / "Ryno Analysis(IBM)"
RYNO_TRANSACTION_PATH = RYNO_DIR / "HI-Small_Trans.csv"
RYNO_TOP20_MANIFEST = RYNO_DIR / "top20" / "manifest.json"
RYNO_BANK12_MANIFEST = RYNO_DIR / "sanitised_bank12" / "manifest.json"
TRANSACTION_TEMPLATE = TEST_DIR / "FrameAML_Transaction_Input_Template.csv"
ACCOUNT_TEMPLATE = TEST_DIR / "FrameAML_Account_Profile_Template.csv"
SAMPLE_INPUT = TEST_DIR / "FrameAML_HI-Small_Synthetic_Development_Fixture.csv"
SAMPLE_LABELS = TEST_DIR / "FrameAML_HI-Small_Synthetic_Fixture_Ground_Truth.csv"

VARIANTS = ["HI-Small", "HI-Medium", "HI-Large", "LI-Small", "LI-Medium", "LI-Large"]
PRIMARY_END = {
    "HI-Small": date(2022, 9, 10),
    "LI-Small": date(2022, 9, 10),
    "HI-Medium": date(2022, 9, 16),
    "LI-Medium": date(2022, 9, 16),
    "HI-Large": date(2022, 11, 5),
    "LI-Large": date(2022, 11, 5),
}
PRIMARY_START = {
    variant: date(2022, 8, 1) if variant.endswith("Large") else date(2022, 9, 1)
    for variant in VARIANTS
}
PATTERN_FILES = [
    "FAN-IN",
    "FAN-OUT",
    "CYCLE",
    "BIPARTITE",
    "GATHER-SCATTER",
    "SCATTER-GATHER",
    "STACK",
    "RANDOM",
]

TRANSACTION_COLUMNS = [
    "Timestamp",
    "From Bank",
    "Account",
    "To Bank",
    "Account.1",
    "Amount Received",
    "Receiving Currency",
    "Amount Paid",
    "Payment Currency",
    "Payment Format",
    "Is Laundering",
]
ACCOUNT_COLUMNS = ["Bank Name", "Bank ID", "Account Number", "Entity ID", "Entity Name"]
HI_SMALL_SOURCE_HEADERS = [
    "Timestamp",
    "From Bank",
    "Account",
    "To Bank",
    "Account",
    "Amount Received",
    "Receiving Currency",
    "Amount Paid",
    "Payment Currency",
    "Payment Format",
    "Is Laundering",
]
EXCEL_MAX_ROWS = 1_048_576
RAW_SHEET_HEADER_ROW = 4
RAW_SHEET_DATA_ROWS = EXCEL_MAX_ROWS - RAW_SHEET_HEADER_ROW
TRANSACTION_SAMPLE_COLUMNS = [
    "Dataset Variant",
    "Transaction Key",
    *TRANSACTION_COLUMNS,
]

NAVY = "17365D"
BLUE = "D9EAF7"
PALE_YELLOW = "FFF2CC"
PALE_GREEN = "E2F0D9"
WHITE = "FFFFFF"
GREY = "595959"
THIN_GREY = Side(style="thin", color="D9E1F2")


class HyperLogLog:
    """Small-memory approximate distinct counter for account-owner entity IDs."""

    def __init__(self, precision: int = 12) -> None:
        self.precision = precision
        self.size = 1 << precision
        self.registers = np.zeros(self.size, dtype=np.uint8)

    def update(self, values: pd.Series) -> None:
        hashes = pd.util.hash_pandas_object(values, index=False).to_numpy(dtype=np.uint64)
        indexes = (hashes >> np.uint64(64 - self.precision)).astype(np.intp)
        remainder = hashes << np.uint64(self.precision)
        nonzero = remainder != 0
        ranks = np.full(len(hashes), 65 - self.precision, dtype=np.uint8)
        if nonzero.any():
            leading = np.floor(
                np.log2(remainder[nonzero].astype(np.float64))
            ).astype(np.int16)
            ranks[nonzero] = (64 - leading).astype(np.uint8)
        np.maximum.at(self.registers, indexes, ranks)

    def estimate(self) -> int:
        alpha = 0.7213 / (1 + 1.079 / self.size)
        inverse_sum = np.sum(np.exp2(-self.registers.astype(np.float64)))
        estimate = alpha * self.size * self.size / inverse_sum
        zeros = int(np.count_nonzero(self.registers == 0))
        if estimate <= 2.5 * self.size and zeros:
            estimate = self.size * math.log(self.size / zeros)
        return int(round(estimate))


def _counter_add(
    counter: Counter, grouped: pd.Series, prefix: tuple[Any, ...] = ()
) -> None:
    for key, value in grouped.items():
        index = key if isinstance(key, tuple) else (key,)
        counter[(*prefix, *index)] += int(value)


def _normalise_bank_id(values: pd.Series) -> pd.Series:
    normalized = values.astype("string").str.lstrip("0")
    return normalized.mask(normalized.eq(""), "0")


def _new_variant_state(manifest: dict[str, Any], variant: str) -> dict[str, Any]:
    tx = manifest["variants"][variant]["transactions"]
    acc = manifest["variants"][variant]["accounts"]
    return {
        "variant": variant,
        "expected_transaction_rows": int(tx["row_count"]),
        "expected_account_rows": int(acc["row_count"]),
        "transaction_rows": 0,
        "positive_transactions": 0,
        "negative_transactions": 0,
        "timestamp_min": None,
        "timestamp_max": None,
        "after_primary_period_rows": 0,
        "after_primary_period_positive_rows": 0,
        "positive_sender_accounts": set(),
        "positive_receiver_accounts": set(),
        "transaction_null_counts": Counter(),
        "account_null_counts": Counter(),
        "account_rows": 0,
        "bank_ids": set(),
        "bank_names": set(),
        "bank_account_counts": Counter(),
        "entity_hll": HyperLogLog(),
        "manifest_transaction_null_counts": tx.get("null_counts", {}),
        "manifest_account_null_counts": acc.get("null_counts", {}),
        "primary_start": PRIMARY_START[variant].isoformat(),
        "primary_end": PRIMARY_END[variant].isoformat(),
        "source_transaction_bytes": int(tx["source_bytes"]),
        "staged_transaction_bytes": (STAGING / f"{variant}_Trans.parquet").stat().st_size,
        "source_account_bytes": int(acc["source_bytes"]),
        "staged_account_bytes": (STAGING / f"{variant}_accounts.parquet").stat().st_size,
    }


def _as_record(row: pd.Series) -> dict[str, Any]:
    record: dict[str, Any] = {}
    for key, value in row.items():
        if pd.isna(value):
            record[key] = None
        elif isinstance(value, pd.Timestamp):
            record[key] = value.isoformat(sep=" ")
        elif isinstance(value, np.generic):
            record[key] = value.item()
        else:
            record[key] = value
    return record


def _amount_bucket_update(
    counter: dict[tuple[str, str, str, int], dict[str, float]],
    frame: pd.DataFrame,
    amount_column: str,
    currency_column: str,
    side: str,
    variant: str,
) -> None:
    work = frame[[currency_column, "Is Laundering", amount_column]].copy()
    work[amount_column] = pd.to_numeric(work[amount_column], errors="coerce")
    work = work.dropna(subset=[currency_column, "Is Laundering", amount_column])
    if work.empty:
        return
    work["_square"] = work[amount_column] * work[amount_column]
    grouped = work.groupby(
        [currency_column, "Is Laundering"], observed=True, sort=False
    )[amount_column].agg(["count", "sum", "min", "max"])
    squares = work.groupby(
        [currency_column, "Is Laundering"], observed=True, sort=False
    )["_square"].sum()
    for (currency, label), values in grouped.iterrows():
        key = (variant, side, str(currency), int(label))
        target = counter.setdefault(
            key,
            {"count": 0, "sum": 0.0, "sum_squares": 0.0, "min": math.inf, "max": -math.inf},
        )
        target["count"] += int(values["count"])
        target["sum"] += float(values["sum"])
        target["sum_squares"] += float(squares.loc[(currency, label)])
        target["min"] = min(target["min"], float(values["min"]))
        target["max"] = max(target["max"], float(values["max"]))


def _update_transaction_profile(
    variant: str,
    parquet_path: Path,
    state: dict[str, Any],
    aggregates: dict[str, Any],
    fixture_candidates: dict[int, list[dict[str, Any]]],
    rng: np.random.Generator,
    batch_size: int,
) -> None:
    parquet = pq.ParquetFile(parquet_path)
    expected_columns = set(TRANSACTION_SAMPLE_COLUMNS)
    if not expected_columns.issubset(set(parquet.schema_arrow.names)):
        raise ValueError(f"{parquet_path.name} is missing expected transaction columns.")

    for batch in parquet.iter_batches(
        batch_size=batch_size, columns=TRANSACTION_SAMPLE_COLUMNS, use_threads=True
    ):
        frame = batch.to_pandas()
        if frame.empty:
            continue
        if not pd.api.types.is_datetime64_any_dtype(frame["Timestamp"]):
            frame["Timestamp"] = pd.to_datetime(frame["Timestamp"], errors="raise")
        state["transaction_rows"] += len(frame)
        state["transaction_null_counts"].update(
            {key: int(value) for key, value in frame.isna().sum().items()}
        )
        labels = frame["Is Laundering"].astype("int8")
        if not labels.isin([0, 1]).all():
            raise ValueError(f"Unexpected Is Laundering value in {parquet_path.name}.")
        positive_count = int(labels.sum())
        state["positive_transactions"] += positive_count
        state["negative_transactions"] += len(frame) - positive_count

        batch_min = frame["Timestamp"].min()
        batch_max = frame["Timestamp"].max()
        state["timestamp_min"] = min(
            state["timestamp_min"], batch_min.isoformat()
        ) if state["timestamp_min"] else batch_min.isoformat()
        state["timestamp_max"] = max(
            state["timestamp_max"], batch_max.isoformat()
        ) if state["timestamp_max"] else batch_max.isoformat()

        beyond = frame["Timestamp"].dt.date > PRIMARY_END[variant]
        state["after_primary_period_rows"] += int(beyond.sum())
        state["after_primary_period_positive_rows"] += int((beyond & labels.eq(1)).sum())

        formats = frame["Payment Format"].fillna("<MISSING>").astype(str)
        _counter_add(
            aggregates["formats"],
            formats.groupby([frame["Payment Format"].fillna("<MISSING>").astype(str), labels])
            .size(),
            prefix=(variant,),
        )
        for side, column in (("paid", "Payment Currency"), ("received", "Receiving Currency")):
            currencies = frame[column].fillna("<MISSING>").astype(str)
            _counter_add(
                aggregates["currencies"],
                pd.DataFrame({"currency": currencies, "label": labels})
                .groupby(["currency", "label"], sort=False)
                .size(),
                prefix=(variant, side),
            )

        pair_frame = pd.DataFrame(
            {
                "payment": frame["Payment Currency"].fillna("<MISSING>").astype(str),
                "receiving": frame["Receiving Currency"].fillna("<MISSING>").astype(str),
                "label": labels.to_numpy(),
            }
        )
        _counter_add(
            aggregates["currency_pairs"],
            pair_frame.groupby(["payment", "receiving", "label"], sort=False).size(),
            prefix=(variant,),
        )

        month_key = frame["Timestamp"].dt.strftime("%Y-%m")
        day_key = frame["Timestamp"].dt.strftime("%Y-%m-%d")
        _counter_add(
            aggregates["monthly"],
            pd.DataFrame({"month": month_key, "label": labels})
            .groupby(["month", "label"], sort=False).size(),
            prefix=(variant,),
        )
        _counter_add(
            aggregates["daily"],
            pd.DataFrame({"day": day_key, "label": labels})
            .groupby(["day", "label"], sort=False).size(),
            prefix=(variant,),
        )
        _counter_add(
            aggregates["primary_period"],
            pd.DataFrame(
                {
                    "period": np.where(beyond, "After stated primary period", "Within stated primary period"),
                    "label": labels,
                }
            ).groupby(["period", "label"], sort=False).size(),
            prefix=(variant,),
        )

        _amount_bucket_update(
            aggregates["amounts"], frame, "Amount Paid", "Payment Currency", "Paid", variant
        )
        _amount_bucket_update(
            aggregates["amounts"], frame, "Amount Received", "Receiving Currency", "Received", variant
        )

        for currency_column, amount_column, side in (
            ("Payment Currency", "Amount Paid", "Paid"),
            ("Receiving Currency", "Amount Received", "Received"),
        ):
            sample_frame = pd.DataFrame(
                {
                    "currency": frame[currency_column].fillna("<MISSING>").astype(str),
                    "label": labels.to_numpy(),
                    "amount": frame[amount_column].to_numpy(),
                }
            )
            for (currency, label), group in sample_frame.groupby(
                ["currency", "label"], observed=True, sort=False
            ):
                take = min(8, len(group))
                indexes = rng.choice(len(group), size=take, replace=False)
                values = group["amount"].to_numpy()[indexes]
                aggregates["amount_samples"][(variant, side, str(currency), int(label))].extend(
                    float(value) for value in values if pd.notna(value)
                )

        positives = frame.loc[labels.eq(1)]
        if not positives.empty:
            sender_keys = (
                variant
                + "|"
                + _normalise_bank_id(positives["From Bank"]).astype(str)
                + "|"
                + positives["Account"].astype(str)
            )
            receiver_keys = (
                variant
                + "|"
                + _normalise_bank_id(positives["To Bank"]).astype(str)
                + "|"
                + positives["Account.1"].astype(str)
            )
            state["positive_sender_accounts"].update(sender_keys.unique())
            state["positive_receiver_accounts"].update(receiver_keys.unique())

        if variant == "HI-Small":
            for label in (0, 1):
                indexes = np.flatnonzero(labels.to_numpy() == label)
                if len(indexes):
                    chosen = rng.choice(indexes, size=min(40, len(indexes)), replace=False)
                    candidates = fixture_candidates[label]
                    candidates.extend(
                        _as_record(frame.iloc[int(index)][TRANSACTION_SAMPLE_COLUMNS])
                        for index in chosen
                    )

    if state["transaction_rows"] != state["expected_transaction_rows"]:
        raise ValueError(
            f"{variant} transaction row mismatch: scanned {state['transaction_rows']:,}, "
            f"manifest says {state['expected_transaction_rows']:,}."
        )


def _update_account_profile(
    variant: str, parquet_path: Path, state: dict[str, Any], batch_size: int
) -> None:
    parquet = pq.ParquetFile(parquet_path)
    required = set(ACCOUNT_COLUMNS)
    if not required.issubset(set(parquet.schema_arrow.names)):
        raise ValueError(f"{parquet_path.name} is missing expected account columns.")
    for batch in parquet.iter_batches(batch_size=batch_size, columns=ACCOUNT_COLUMNS, use_threads=True):
        frame = batch.to_pandas()
        state["account_rows"] += len(frame)
        state["account_null_counts"].update(
            {key: int(value) for key, value in frame.isna().sum().items()}
        )
        state["entity_hll"].update(frame["Entity ID"])
        state["bank_ids"].update(frame["Bank ID"].dropna().astype(str).unique())
        state["bank_names"].update(frame["Bank Name"].dropna().astype(str).unique())
        counts = frame.groupby(["Bank ID", "Bank Name"], dropna=False, sort=False).size()
        for (bank_id, bank_name), count in counts.items():
            state["bank_account_counts"][(str(bank_id), str(bank_name))] += int(count)
    if state["account_rows"] != state["expected_account_rows"]:
        raise ValueError(
            f"{variant} account row mismatch: scanned {state['account_rows']:,}, "
            f"manifest says {state['expected_account_rows']:,}."
        )


def _parse_patterns(variant: str, path: Path) -> dict[str, dict[str, int]]:
    stats: dict[str, dict[str, int]] = defaultdict(lambda: {"attempts": 0, "annotation_lines": 0})
    active_family: str | None = None
    begin = re.compile(r"^BEGIN LAUNDERING ATTEMPT - ([A-Z-]+)")
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        match = begin.match(line.strip())
        if match:
            active_family = match.group(1)
            stats[active_family]["attempts"] += 1
        elif line.startswith("END LAUNDERING ATTEMPT"):
            active_family = None
        elif active_family and line.strip():
            stats[active_family]["annotation_lines"] += 1
    return {family: dict(values) for family, values in stats.items()}


def _profile_all(variants: list[str], batch_size: int) -> dict[str, Any]:
    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    if manifest.get("status") != "complete":
        raise ValueError("Staging manifest is not marked complete.")
    aggregates: dict[str, Any] = {
        "formats": Counter(),
        "currencies": Counter(),
        "currency_pairs": Counter(),
        "monthly": Counter(),
        "daily": Counter(),
        "primary_period": Counter(),
        "amounts": {},
        "amount_samples": defaultdict(list),
    }
    rng = np.random.default_rng(20261026)
    fixture_candidates: dict[int, list[dict[str, Any]]] = {0: [], 1: []}
    states: dict[str, dict[str, Any]] = {}
    patterns: dict[str, dict[str, dict[str, int]]] = {}
    for variant in variants:
        print(f"Scanning every transaction row: {variant}", flush=True)
        state = _new_variant_state(manifest, variant)
        _update_transaction_profile(
            variant,
            STAGING / f"{variant}_Trans.parquet",
            state,
            aggregates,
            fixture_candidates,
            rng,
            batch_size,
        )
        print(f"Scanning every account record: {variant}", flush=True)
        _update_account_profile(
            variant,
            STAGING / f"{variant}_accounts.parquet",
            state,
            batch_size,
        )
        pattern_path = STAGING / "patterns" / f"{variant}_Patterns.txt"
        patterns[variant] = _parse_patterns(variant, pattern_path)
        state["approximate_unique_entity_ids"] = state["entity_hll"].estimate()
        state["unique_bank_id_source_values"] = len(state["bank_ids"])
        state["unique_bank_names"] = len(state["bank_names"])
        state["positive_sender_account_count"] = len(state["positive_sender_accounts"])
        state["positive_receiver_account_count"] = len(state["positive_receiver_accounts"])
        state["positive_endpoint_account_count"] = len(
            state["positive_sender_accounts"] | state["positive_receiver_accounts"]
        )
        state["bank_account_counts"] = [
            {"bank_id": bank_id, "bank_name": bank_name, "account_rows": count}
            for (bank_id, bank_name), count in sorted(state["bank_account_counts"].items())
        ]
        state["transaction_null_counts"] = dict(state["transaction_null_counts"])
        state["account_null_counts"] = dict(state["account_null_counts"])
        for key in (
            "positive_sender_accounts",
            "positive_receiver_accounts",
            "bank_ids",
            "bank_names",
            "entity_hll",
        ):
            state.pop(key)
        state["account_composite_key_unique_per_staging_validation"] = bool(
            manifest["variants"][variant]["accounts"].get("composite_join_key_unique")
        )
        state["source_transaction_columns"] = manifest["variants"][variant]["transactions"][
            "source_columns"
        ]
        state["staged_transaction_columns"] = manifest["variants"][variant]["transactions"][
            "columns"
        ]
        state["source_account_columns"] = manifest["variants"][variant]["accounts"][
            "source_columns"
        ]
        state["pattern_file_bytes"] = int(
            manifest["variants"][variant]["patterns"]["source_bytes"]
        )
        states[variant] = state
        print(
            f"Completed {variant}: {state['transaction_rows']:,} transactions, "
            f"{state['account_rows']:,} account rows, {state['positive_transactions']:,} "
            "positive labels.",
            flush=True,
        )

    result = {
        "profiled_at": "2026-10-26",
        "source": {
            "kaggle_url": "https://www.kaggle.com/datasets/ealtman2019/ibm-transactions-for-anti-money-laundering-aml",
            "paper_url": "https://arxiv.org/abs/2306.16424",
            "manifest": str(MANIFEST_PATH),
            "variants_profiled": variants,
            "full_six_variant_profile": set(variants) == set(VARIANTS),
            "rows_are_full_population": True,
            "amount_percentiles_are_sampled": True,
            "amount_sample_method": (
                "Deterministic random sample within each row batch, stratified by "
                "variant, amount side, currency and benchmark label; exact count, "
                "sum, minimum and maximum are calculated over all rows."
            ),
        },
        "variants": states,
        "aggregates": {
            "formats": [
                {"variant": key[0], "format": key[1], "label": key[2], "count": value}
                for key, value in aggregates["formats"].items()
            ],
            "currencies": [
                {"variant": key[0], "side": key[1], "currency": key[2], "label": key[3], "count": value}
                for key, value in aggregates["currencies"].items()
            ],
            "currency_pairs": [
                {"variant": key[0], "payment_currency": key[1], "receiving_currency": key[2], "label": key[3], "count": value}
                for key, value in aggregates["currency_pairs"].items()
            ],
            "monthly": [
                {"variant": key[0], "month": key[1], "label": key[2], "count": value}
                for key, value in aggregates["monthly"].items()
            ],
            "daily": [
                {"variant": key[0], "day": key[1], "label": key[2], "count": value}
                for key, value in aggregates["daily"].items()
            ],
            "primary_period": [
                {"variant": key[0], "period": key[1], "label": key[2], "count": value}
                for key, value in aggregates["primary_period"].items()
            ],
            "amounts": [
                {
                    "variant": key[0],
                    "side": key[1],
                    "currency": key[2],
                    "label": key[3],
                    **values,
                    "approximate_p05": float(np.percentile(aggregates["amount_samples"][key], 5))
                    if aggregates["amount_samples"][key] else None,
                    "approximate_median": float(np.percentile(aggregates["amount_samples"][key], 50))
                    if aggregates["amount_samples"][key] else None,
                    "approximate_p95": float(np.percentile(aggregates["amount_samples"][key], 95))
                    if aggregates["amount_samples"][key] else None,
                    "percentile_sample_n": len(aggregates["amount_samples"][key]),
                }
                for key, values in aggregates["amounts"].items()
            ],
        },
        "pattern_annotations": patterns,
        "fixture_candidates": fixture_candidates,
        "staging_totals": {
            "transaction_rows": sum(state["transaction_rows"] for state in states.values()),
            "account_rows": sum(state["account_rows"] for state in states.values()),
            "positive_transactions": sum(state["positive_transactions"] for state in states.values()),
            "variant_count": len(states),
        },
    }
    PROFILE_PATH.parent.mkdir(parents=True, exist_ok=True)
    PROFILE_PATH.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    return result


def _table_sheet(
    workbook: Workbook,
    title: str,
    subtitle: str,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[float],
    name: str,
) -> Any:
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "4472C4"
    end_column = get_column_letter(len(headers))
    sheet.merge_cells(f"A1:{end_column}1")
    top = sheet["A1"]
    top.value = title
    top.fill = PatternFill("solid", fgColor=NAVY)
    top.font = Font(name="Aptos Display", size=17, bold=True, color=WHITE)
    top.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(f"A2:{end_column}2")
    note = sheet["A2"]
    note.value = subtitle
    note.font = Font(name="Aptos", size=9, italic=True, color=GREY)
    note.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 32
    for index, header in enumerate(headers, 1):
        cell = sheet.cell(4, index, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN_GREY)
        sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
    sheet.row_dimensions[4].height = 32
    body_font = Font(name="Aptos", size=8.5, color="202124")
    body_alignment = Alignment(wrap_text=True, vertical="top")
    body_border = Border(bottom=THIN_GREY)
    stripe_fill = PatternFill("solid", fgColor="F7F9FC")
    if len(rows) > 10_000:
        for values in rows:
            sheet.append(values)
    else:
        for row_number, values in enumerate(rows, 5):
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row_number, column, value)
                cell.font = body_font
                cell.alignment = body_alignment
                cell.border = body_border
                if isinstance(value, float):
                    cell.number_format = "#,##0.00;[Red](#,##0.00);-"
            if row_number % 2:
                for cell in sheet[row_number]:
                    cell.fill = stripe_fill
    if rows:
        ref = f"A4:{end_column}{4 + len(rows)}"
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
        sheet.auto_filter.ref = ref
    sheet.freeze_panes = "A5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"
    sheet.sheet_properties.outlinePr.summaryBelow = True
    return sheet


def _raw_sheet_ranges(transaction_rows: int) -> list[tuple[int, int]]:
    return [
        (
            start,
            min(start + RAW_SHEET_DATA_ROWS - 1, transaction_rows),
        )
        for start in range(1, transaction_rows + 1, RAW_SHEET_DATA_ROWS)
    ]


def _add_hismall_raw_sheets(
    workbook: Workbook,
    transaction_path: Path,
    expected_rows: int,
) -> None:
    if not transaction_path.is_file():
        raise FileNotFoundError(
            f"Cannot add the full HI-Small source tabs; source CSV is missing: "
            f"{transaction_path}"
        )

    ranges = _raw_sheet_ranges(expected_rows)
    reader = None
    with transaction_path.open("r", encoding="utf-8-sig", newline="") as source_file:
        reader = csv.reader(source_file)
        try:
            headers = next(reader)
        except StopIteration as exc:
            raise ValueError(f"HI-Small source CSV is empty: {transaction_path}") from exc
        if headers != HI_SMALL_SOURCE_HEADERS:
            raise ValueError(
                "HI-Small source columns or order do not match the expected source "
                f"schema: {headers}"
            )

        source_row_number = 0
        for sheet_number, (first_source_row, last_source_row) in enumerate(ranges, 1):
            sheet = workbook.create_sheet(f"HI-Small Raw {sheet_number}")
            sheet.sheet_view.showGridLines = True
            sheet.sheet_properties.tabColor = "70AD47"
            sheet.freeze_panes = f"A{RAW_SHEET_HEADER_ROW + 1}"
            sheet.sheet_view.zoomScale = 85
            widths = [21, 14, 18, 14, 18, 18, 18, 18, 18, 18, 16]
            for column, width in enumerate(widths, 1):
                sheet.column_dimensions[get_column_letter(column)].width = width

            title = WriteOnlyCell(
                sheet, value=f"HI-Small full source | rows {first_source_row:,}-{last_source_row:,}"
            )
            title.fill = PatternFill("solid", fgColor=NAVY)
            title.font = Font(name="Aptos Display", size=16, bold=True, color=WHITE)
            title.alignment = Alignment(vertical="center")
            sheet.append([title] + [None] * (len(headers) - 1))
            sheet.row_dimensions[1].height = 30

            subtitle = WriteOnlyCell(
                sheet,
                value=(
                    "Original HI-Small transaction CSV rows in source order. Values are "
                    "preserved as text for exact comparison; the source repeats the "
                    "'Account' header for the receiver account."
                ),
            )
            subtitle.font = Font(name="Aptos", size=9, italic=True, color=GREY)
            subtitle.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.append([subtitle] + [None] * (len(headers) - 1))
            sheet.row_dimensions[2].height = 32
            sheet.append([None] * len(headers))
            header_cells = [
                WriteOnlyCell(sheet, value=value) for value in headers
            ]
            for cell in header_cells:
                cell.fill = PatternFill("solid", fgColor=NAVY)
                cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.append(header_cells)
            sheet.row_dimensions[RAW_SHEET_HEADER_ROW].height = 34

            expected_sheet_rows = last_source_row - first_source_row + 1
            for _ in range(expected_sheet_rows):
                try:
                    row = next(reader)
                except StopIteration as exc:
                    raise ValueError(
                        f"HI-Small source ended at row {source_row_number:,}; expected "
                        f"{expected_rows:,} rows."
                    ) from exc
                source_row_number += 1
                if len(row) != len(headers):
                    raise ValueError(
                        f"HI-Small source row {source_row_number:,} has "
                        f"{len(row)} fields; expected {len(headers)}."
                    )
                sheet.append(row)

            sheet.auto_filter.ref = (
                f"A{RAW_SHEET_HEADER_ROW}:"
                f"{get_column_letter(len(headers))}"
                f"{RAW_SHEET_HEADER_ROW + expected_sheet_rows}"
            )

        try:
            next(reader)
        except StopIteration:
            pass
        else:
            raise ValueError(
                f"HI-Small source contains more than the expected {expected_rows:,} rows."
            )


def _rewrite_excel_compatible(
    source_path: Path,
    target_path: Path,
    include_hismall_raw: bool,
    expected_hismall_rows: int,
    metadata_workbook: Workbook,
) -> None:
    source = load_workbook(source_path, read_only=True, data_only=False)
    output = Workbook(write_only=True)
    required_sheets = {
        "Start Here",
        "CSV Mapping",
        "Account Type Guide",
        "Bank Identifier Guide",
        "Ryno Comparison",
        "Ryno Field Mapping",
    }
    if not required_sheets.issubset(source.sheetnames):
        source.close()
        raise ValueError("The source workbook is missing required sheets.")

    try:
        for sheet_name in source.sheetnames:
            source_sheet = source[sheet_name]
            metadata_sheet = metadata_workbook[sheet_name]
            output_sheet = output.create_sheet(sheet_name)
            output_sheet.sheet_view.showGridLines = False
            output_sheet.sheet_view.zoomScale = metadata_sheet.sheet_view.zoomScale
            output_sheet.freeze_panes = metadata_sheet.freeze_panes or "A5"
            if metadata_sheet.sheet_properties.tabColor:
                output_sheet.sheet_properties.tabColor = copy(
                    metadata_sheet.sheet_properties.tabColor
                )
            for column in range(1, source_sheet.max_column + 1):
                letter = get_column_letter(column)
                source_dimension = metadata_sheet.column_dimensions.get(letter)
                if source_dimension and source_dimension.width:
                    output_sheet.column_dimensions[letter].width = source_dimension.width
                else:
                    output_sheet.column_dimensions[letter].width = 22

            for row_number, row in enumerate(
                source_sheet.iter_rows(values_only=False), 1
            ):
                output_row = []
                for cell in row:
                    has_style = getattr(cell, "has_style", False)
                    if cell.value is None and not has_style:
                        output_row.append(None)
                        continue
                    output_cell = WriteOnlyCell(output_sheet, value=cell.value)
                    if has_style:
                        output_cell.font = copy(cell.font)
                        output_cell.fill = copy(cell.fill)
                        output_cell.border = copy(cell.border)
                        output_cell.alignment = copy(cell.alignment)
                        output_cell.number_format = cell.number_format
                        output_cell.protection = copy(cell.protection)
                    hyperlink = getattr(cell, "hyperlink", None)
                    if hyperlink:
                        output_cell.hyperlink = hyperlink.target
                    output_row.append(output_cell)
                if row_number == 1 and output_row and output_row[0] is not None:
                    title_style = output_row[0]._style
                    for column in range(1, source_sheet.max_column):
                        if output_row[column] is None:
                            output_row[column] = WriteOnlyCell(output_sheet)
                        output_row[column]._style = title_style
                output_sheet.append(output_row)
                dimension = metadata_sheet.row_dimensions.get(row_number)
                if dimension and dimension.height:
                    output_sheet.row_dimensions[row_number].height = dimension.height

            for table in metadata_sheet.tables.values():
                headers = [
                    str(metadata_sheet.cell(4, column).value)
                    for column in range(1, source_sheet.max_column + 1)
                ]
                if len(headers) != len(set(headers)):
                    raise ValueError(
                        f"Cannot preserve Excel table {table.displayName}; "
                        "its header names are not unique."
                    )
                output_table = Table(
                    displayName=table.displayName,
                    ref=table.ref,
                    tableColumns=[
                        TableColumn(id=index, name=header)
                        for index, header in enumerate(headers, 1)
                    ],
                    autoFilter=AutoFilter(ref=table.ref),
                    tableStyleInfo=TableStyleInfo(
                        name=(
                            table.tableStyleInfo.name
                            if table.tableStyleInfo
                            else "TableStyleMedium2"
                        ),
                        showFirstColumn=(
                            table.tableStyleInfo.showFirstColumn
                            if table.tableStyleInfo
                            else False
                        ),
                        showLastColumn=(
                            table.tableStyleInfo.showLastColumn
                            if table.tableStyleInfo
                            else False
                        ),
                        showRowStripes=(
                            table.tableStyleInfo.showRowStripes
                            if table.tableStyleInfo
                            else True
                        ),
                        showColumnStripes=(
                            table.tableStyleInfo.showColumnStripes
                            if table.tableStyleInfo
                            else False
                        ),
                    ),
                )
                output_sheet.add_table(output_table)
            if metadata_sheet.auto_filter.ref and not metadata_sheet.tables:
                output_sheet.auto_filter.ref = metadata_sheet.auto_filter.ref
            if sheet_name == "Variant Overview" and source_sheet.max_row >= 5:
                output_sheet.conditional_formatting.add(
                    f"E5:E{source_sheet.max_row}",
                    ColorScaleRule(
                        start_type="min",
                        start_color="E2F0D9",
                        end_type="max",
                        end_color="F4CCCC",
                    ),
                )

        if include_hismall_raw:
            _add_hismall_raw_sheets(
                output,
                RYNO_TRANSACTION_PATH,
                expected_hismall_rows,
            )

        output.save(target_path)
    finally:
        source.close()


def _count_by_variant_label(
    rows: list[dict[str, Any]], group_field: str, value_field: str
) -> dict[tuple[str, str], dict[int, int]]:
    result: dict[tuple[str, str], dict[int, int]] = defaultdict(lambda: {0: 0, 1: 0})
    for row in rows:
        result[(row["variant"], str(row[group_field]))][int(row["label"])] += int(row[value_field])
    return result


def _select_fixture(profile: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    candidate_map = profile["fixture_candidates"]
    rng = np.random.default_rng(20261026)
    selected: dict[int, list[dict[str, Any]]] = {}
    for label in (0, 1):
        candidates = candidate_map.get(str(label), candidate_map.get(label, []))
        take = min(100, len(candidates))
        if take == 0:
            raise ValueError(f"HI-Small fixture has no label={label} candidate rows.")
        indexes = rng.choice(len(candidates), size=take, replace=False)
        selected[label] = [candidates[int(index)] for index in indexes]
    transactions = selected[0] + selected[1]
    return transactions, [
        {
            "Dataset Variant": record["Dataset Variant"],
            "Unique Identifier for the transaction": record["Transaction Key"],
            "IBM Benchmark Label (evaluation only)": int(record["Is Laundering"]),
        }
        for record in transactions
    ]


def _build_csv_files(profile: dict[str, Any]) -> None:
    TEST_DIR.mkdir(parents=True, exist_ok=True)
    transaction_columns = [
        "Dataset Variant (benchmark only)",
        "Account Key",
        "Transaction Amount in Account Currency",
        "Transaction Amount in Original Currency",
        "Original Currency Code",
        "Transaction Branch",
        "Transaction DateTime",
        "Opposite Account Number",
        "Opposite Bank",
        "Opposite Country Code",
        "Transaction Type",
        "Transaction Direction",
        "Unique Identifier for the transaction",
    ]
    account_columns = [
        "Dataset Variant (benchmark only)",
        "Account Key",
        "Bank ID",
        "Account Number",
        "Bank Name",
        "Entity ID",
        "Entity Name",
        "Account Type",
        "Account Currency",
        "Branch",
        "Country Code",
        "Effective From",
        "Effective To",
    ]
    pd.DataFrame(columns=transaction_columns).to_csv(TRANSACTION_TEMPLATE, index=False, encoding="utf-8")
    pd.DataFrame(columns=account_columns).to_csv(ACCOUNT_TEMPLATE, index=False, encoding="utf-8")

    transactions, labels = _select_fixture(profile)
    events: list[dict[str, Any]] = []
    for record in transactions:
        variant = record["Dataset Variant"]
        timestamp = record["Timestamp"]
        label = int(record["Is Laundering"])
        transaction_key = record["Transaction Key"]
        from_bank = str(record["From Bank"])
        to_bank = str(record["To Bank"])
        sender_account = str(record["Account"])
        receiver_account = str(record["Account.1"])
        payment_amount = float(record["Amount Paid"])
        received_amount = float(record["Amount Received"])

        sender_key = f"{variant}|{int(from_bank)}|{sender_account}"
        receiver_key = f"{variant}|{int(to_bank)}|{receiver_account}"
        shared = {
            "Dataset Variant (benchmark only)": variant,
            "Transaction Branch": "",
            "Opposite Country Code": "",
            "Unique Identifier for the transaction": transaction_key,
        }
        events.append(
            {
                **shared,
                "Account Key": sender_key,
                "Transaction Amount in Account Currency": "",
                "Transaction Amount in Original Currency": payment_amount,
                "Original Currency Code": record["Payment Currency"],
                "Transaction DateTime": timestamp,
                "Opposite Account Number": receiver_account,
                "Opposite Bank": to_bank,
                "Transaction Type": record["Payment Format"],
                "Transaction Direction": "Debit",
            }
        )
        events.append(
            {
                **shared,
                "Account Key": receiver_key,
                "Transaction Amount in Account Currency": "",
                "Transaction Amount in Original Currency": received_amount,
                "Original Currency Code": record["Receiving Currency"],
                "Transaction DateTime": timestamp,
                "Opposite Account Number": sender_account,
                "Opposite Bank": from_bank,
                "Transaction Type": record["Payment Format"],
                "Transaction Direction": "Credit",
            }
        )
    pd.DataFrame(events, columns=transaction_columns).to_csv(
        SAMPLE_INPUT, index=False, encoding="utf-8"
    )
    pd.DataFrame(labels).to_csv(SAMPLE_LABELS, index=False, encoding="utf-8")


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source_file:
        for block in iter(lambda: source_file.read(8 * 1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_required_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"Required Ryno provenance manifest is missing: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def _ryno_comparison_rows(
    profile: dict[str, Any],
    source_sha256: str,
) -> list[list[Any]]:
    top20 = _read_required_json(RYNO_TOP20_MANIFEST)
    bank12 = _read_required_json(RYNO_BANK12_MANIFEST)
    top20_hash = top20["inputs"]["HI-Small_Trans.csv"]
    bank12_hash = bank12["inputs"]["HI-Small_Trans.csv"]
    if top20_hash != bank12_hash or source_sha256 != top20_hash:
        raise ValueError(
            "The HI-Small source CSV does not match the SHA-256 recorded by both "
            "Ryno conversion manifests."
        )

    top20_counts = top20["counts"]
    bank12_counts = bank12["counts"]
    top20_parameters = top20["parameters"]
    bank12_parameters = bank12["parameters"]
    source_state = profile["variants"]["HI-Small"]

    return [
        [
            "Original HI-Small source",
            "Transaction rows",
            int(source_state["transaction_rows"]),
            "Payments",
            "One row is one directed source payment. The original source has 11 columns.",
        ],
        [
            "Original HI-Small source",
            "Label-positive rows",
            int(source_state["positive_transactions"]),
            "Payments",
            "Synthetic Is Laundering label; not evidence of real-world wrongdoing.",
        ],
        [
            "Original HI-Small source",
            "Account-dimension rows",
            int(source_state["account_rows"]),
            "Accounts",
            "Separate HI-Small_accounts.csv dimension; not repeated as transaction legs.",
        ],
        [
            "Source identity",
            "SHA-256",
            source_sha256,
            "Checksum",
            "Matches the HI-Small transaction input hash in both Ryno manifests.",
        ],
        [
            "Ryno top-20 conversion",
            "Included banks",
            len(top20["banks"]),
            "Banks",
            f"Bank IDs: {', '.join(str(bank) for bank in top20['banks'])}.",
        ],
        [
            "Ryno top-20 conversion",
            "Excluded bank",
            top20_parameters.get("exclude_banks", ""),
            "Bank ID",
            "Hub bank exclusion is part of the selected conversion scope.",
        ],
        [
            "Ryno top-20 conversion",
            "Payments retained",
            int(top20_counts["payments"]),
            "Payments",
            "Tail excluded by the recorded parameters; this is a subset of full HI-Small.",
        ],
        [
            "Ryno top-20 conversion",
            "Account-side message legs",
            int(top20_counts["legs"]),
            "Debit/credit legs",
            "A source payment may create a debit, a credit, or two in-scope account legs.",
        ],
        [
            "Ryno top-20 conversion",
            "Positive payments",
            int(top20_counts["pos_payments"]),
            "Payments",
            f"{int(top20_counts['pos_legs']):,} positive account-side legs; labels are separate from messages.",
        ],
        [
            "Ryno top-20 conversion",
            "Positive payments in pattern annotations",
            int(top20_counts["pos_payments_in_patterns"]),
            "Payments",
            "Pattern text does not annotate every label-positive payment.",
        ],
        [
            "Ryno top-20 conversion",
            "Tail rows dropped",
            int(top20_counts.get("dropped", {}).get("tail", 0)),
            "Source rows",
            "Tail exclusion differs from the Bank-12 conversion, which retains the tail.",
        ],
        [
            "Ryno Bank-12 sanitised conversion",
            "Included bank",
            ", ".join(str(bank) for bank in bank12["banks"]),
            "Bank ID",
            "Single-bank view; not the full HI-Small population.",
        ],
        [
            "Ryno Bank-12 sanitised conversion",
            "Payments represented",
            int(bank12_counts["payments"]),
            "Payments",
            "Selected from payments involving Bank 12; not a random or complete benchmark slice.",
        ],
        [
            "Ryno Bank-12 sanitised conversion",
            "Account-side output rows",
            int(bank12_counts["legs"]),
            "Debit/credit legs",
            f"{int(bank12_counts['legs_D']):,} debit and {int(bank12_counts['legs_C']):,} credit legs.",
        ],
        [
            "Ryno Bank-12 sanitised conversion",
            "Positive payments",
            int(bank12_counts["pos_payments"]),
            "Payments",
            f"{int(bank12_counts['pos_legs']):,} positive legs; labels are provided separately.",
        ],
        [
            "Ryno Bank-12 assumptions",
            "Account currency",
            bank12.get("accountCurrency", ""),
            "Synthetic assumption",
            "All accounts are assigned ZAR for the test; HI-Small does not provide account currency.",
        ],
        [
            "Ryno Bank-12 assumptions",
            "ZAR/USD conversion",
            bank12_parameters.get("usd_rate", ""),
            "Approximate rate",
            "The supplied note describes 17.50 as an approximation, not an official rate.",
        ],
        [
            "Ryno Bank-12 assumptions",
            "Synthetic cross-border legs",
            int(bank12_counts["cross_border_legs"]),
            "Debit/credit legs",
            "Derived using bank-name country mapping; Bank 12 is treated as a US bank.",
        ],
        [
            "Ryno Bank-12 assumptions",
            "Synthetic teller legs",
            int(bank12_counts["teller_legs"]),
            "Debit/credit legs",
            "Branch is made up; cash/cheque use fixed home branch, electronic legs use virtual branch.",
        ],
        [
            "Ryno Notes",
            "Multicurrency pattern annotations",
            "216 / 370",
            "Patterns",
            "Count cited in Ryno's notes; not independently recomputed in this workbook.",
        ],
        [
            "Ryno Notes",
            "Multicurrency cycles",
            "38 / 54",
            "Cycles",
            "Count cited in Ryno's notes; illustrates why account-leg-only benchmarking loses some relations.",
        ],
        [
            "Interpretation",
            "Benchmark comparability",
            "Not like-for-like",
            "Scope warning",
            "Bank filtering, date-tail rules, labels, payment-versus-leg units, and synthetic fields differ. Do not compare detection performance as if these were identical populations.",
        ],
        [
            "Interpretation",
            "UDM representation limitation",
            "Relations can be implicit or lost",
            "Modeling caution",
            "Ryno's notes describe ordered chains, cycles, gather/scatter structures, shared counterparties, and cross-currency hops. Test these as named scenarios rather than assuming a UDM-only feed preserves them.",
        ],
    ]


def _build_workbook(
    profile: dict[str, Any],
    include_hismall_raw: bool = True,
) -> None:
    if "HI-Small" not in profile["variants"]:
        raise ValueError(
            "The workbook requires the HI-Small variant for its Ryno comparison. "
            "Include HI-Small in --variants."
        )
    hismall_rows = int(profile["variants"]["HI-Small"]["transaction_rows"])
    output_path = WORKBOOK_PATH if include_hismall_raw else SUMMARY_WORKBOOK_PATH
    workbook = Workbook()
    start = workbook.active
    start.title = "Start Here"
    start.sheet_view.showGridLines = False
    start.sheet_view.zoomScale = 90
    start.sheet_properties.tabColor = NAVY
    start.column_dimensions["A"].width = 31
    start.column_dimensions["B"].width = 96
    start.merge_cells("A1:B1")
    start["A1"] = "FrameAML | IBM AML Data Preparation & Descriptive Analysis"
    start["A1"].fill = PatternFill("solid", fgColor=NAVY)
    start["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    start["A1"].alignment = Alignment(vertical="center")
    start.row_dimensions[1].height = 34
    start.merge_cells("A2:B2")
    start["A2"] = (
        "Descriptive profile of six synthetic IBM variants, full HI-Small source "
        "transactions, and Ryno's UDM-style transformations. Use the sheet guide "
        "below to navigate the workbook."
    )
    start["A2"].font = Font(name="Aptos", size=10, italic=True, color=GREY)
    start["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    start.row_dimensions[2].height = 34
    start_values = [
        ("Analysis scope", "This is descriptive analysis only. It reports population counts, distributions, dates, data quality, source mappings, and limitations; it does not claim a detection model or prove real-world laundering."),
        ("Six-variant population", f"{profile['staging_totals']['transaction_rows']:,} transaction rows and {profile['staging_totals']['account_rows']:,} account rows across {profile['staging_totals']['variant_count']} IBM variants. Variants are separate alternatives; do not combine them into one time series without retaining the variant."),
        ("HI-Small full source", f"{hismall_rows:,} original HI-Small transactions and {profile['variants']['HI-Small']['account_rows']:,} account rows. " + ("All transaction rows appear in the five HI-Small Raw tabs, split only because Excel allows at most 1,048,576 rows per worksheet." if include_hismall_raw else "This summary-only version does not embed raw rows; use the standard build to include the five complete HI-Small Raw tabs.")),
        ("Raw data fidelity", "The five HI-Small Raw tabs preserve the original 11 columns, their order, repeated Account header, source row order, and source values as text. Power BI staging adds derived fields such as Dataset Variant, Transaction Key, sender/receiver aliases, and debit/credit aliases; those are not original CSV columns. Source row ranges are stated on each tab; the raw HI-Small CSV remains the immutable source of truth." if include_hismall_raw else "Raw transaction rows are not embedded in this summary-only workbook; the source CSV remains unchanged in Ryno Analysis(IBM). The Power BI staging adds derived identifiers and sender/receiver/debit/credit aliases."),
        ("Ryno comparison", "Ryno's top-20-bank and Bank-12 outputs are account-side UDM message legs, not a second copy of all HI-Small payments. They use different bank scopes, filtering, label coverage, and assumed fields, so their counts and labels are not a like-for-like benchmark."),
        ("Ryno field mapping", "The Ryno Field Mapping sheet lists all 12 requested attributes in order and identifies direct source mappings, derived values, and assumptions that need confirmation before system use."),
        ("Where source data lives", "Original six-variant CSVs are in archive\\; staged Parquet tables are in powerbi_staging\\; Ryno's HI-Small source, scripts, notes, and conversion manifests are in Ryno Analysis(IBM)\\. Large source/generated data are kept locally and excluded from normal Git commits."),
        ("Excel limits", "A worksheet holds at most 1,048,576 rows including its title and header rows. HI-Small is therefore split across five raw-data tabs; the five tabs together contain all 5,078,345 transactions."),
        ("The supplied PBIX", "Power BI files are report/model artifacts, not proof of full source coverage. The complete population is validated against CSV row counts and staged Parquet; the workbook contains the complete HI-Small transaction source, not all six variants' raw data."),
        ("Account types", "There is no Account Type, account product, account currency, branch, or country-code column in the IBM transaction/account schemas. The account file has Bank Name, Bank ID, Account Number, Entity ID and Entity Name. The Account Type Guide defines common roles such as clearing and settlement but makes clear these cannot be assigned from this data alone; leave missing target fields blank pending trusted account-master evidence."),
        ("Bank identification", "Use the IBM account dimension's own Bank Name joined by variant, normalized bank ID and account number. IBM Bank IDs are synthetic benchmark-local keys, not real worldwide routing codes. See Bank Identifier Guide for BIC, LEI, national-code distinctions and safe enrichment design."),
        ("Benchmark meaning", "Is Laundering is an IBM synthetic transaction label. It is suitable as benchmark truth within the simulator, not proof about a real person or payment. Keep it out of the scoring input and join the separate fixture ground-truth file only after scoring."),
        ("Account-level label", "Account labels are not directly supplied as a separate ground-truth field. Defining an account as positive if it ever appears in a positive transaction is a derived roll-up; it is not equivalent to transaction-level detection and can leak future outcomes if used without time-aware rules."),
        ("Currency rule", "Amount Paid belongs with Payment Currency. Amount Received belongs with Receiving Currency. Account currency is not supplied, so Transaction Amount in Account Currency remains blank until an external account-currency source or an explicit scenario assumption is provided."),
        ("Recommended reading order", "Start Here → Variant Overview → Ryno Comparison → " + ("HI-Small Raw 1–5 → " if include_hismall_raw else "") + "Field Dictionary → CSV Mapping → Account Profile → Account Type Guide → Bank Identifier Guide → Pattern Catalogue → Test Protocol. Use the sheet-by-sheet guide below for every tab."),
        ("Profile reproducibility", "Run .venv\\Scripts\\python.exe analysis\\build_frame_aml_descriptive_workbook.py to scan staged Parquet and build this workbook. The full HI-Small CSV and Ryno manifests must be present. Use --reuse-profile to skip the population scan; use --without-hi-small-raw for a smaller summary-only workbook."),
    ]
    for row_index, (label, text) in enumerate(start_values, 4):
        start.cell(row_index, 1, label)
        start.cell(row_index, 2, text)
        start.cell(row_index, 1).font = Font(name="Aptos", size=9, bold=True, color=NAVY)
        start.cell(row_index, 1).fill = PatternFill("solid", fgColor=BLUE)
        start.cell(row_index, 2).fill = PatternFill(
            "solid", fgColor="F7F9FC" if row_index % 2 == 0 else WHITE
        )
        for cell in start[row_index]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN_GREY)
        start.row_dimensions[row_index].height = 58
    start.freeze_panes = "A4"

    variant_rows: list[list[Any]] = []
    for variant in VARIANTS:
        if variant not in profile["variants"]:
            continue
        state = profile["variants"][variant]
        total = state["transaction_rows"]
        positives = state["positive_transactions"]
        rate = positives / total if total else 0
        variant_rows.append(
            [
                variant,
                total,
                state["negative_transactions"],
                positives,
                rate,
                state["account_rows"],
                state["approximate_unique_entity_ids"],
                "Approx. HLL; see note",
                state["unique_bank_id_source_values"],
                state["unique_bank_names"],
                state["timestamp_min"],
                state["timestamp_max"],
                f"{state['primary_start']} to {state['primary_end']}",
                state["after_primary_period_rows"],
                state["after_primary_period_positive_rows"],
                state["source_transaction_bytes"],
                state["staged_transaction_bytes"],
            ]
        )
    overview = _table_sheet(
        workbook,
        "Variant Overview",
        "Exact full-population counts from all six staged variants. Dates after the Kaggle-listed primary period are shown separately; they are an essential split/leakage check.",
        [
            "Dataset Variant", "Transactions", "Label 0", "Label 1", "Positive Rate",
            "Account Rows", "Approx. Unique Entity IDs", "Entity Count Method",
            "Distinct Bank ID Values", "Distinct Bank Names", "Actual Min Timestamp",
            "Actual Max Timestamp", "Kaggle Primary Period", "Rows After Primary End",
            "Positive Rows After Primary End", "Source CSV Bytes", "Staged Parquet Bytes",
        ],
        variant_rows,
        [15, 16, 15, 15, 14, 15, 22, 20, 18, 17, 22, 22, 27, 19, 23, 19, 21],
        "VariantOverviewTable",
    )
    overview.conditional_formatting.add(
        f"E5:E{4 + len(variant_rows)}",
        ColorScaleRule(start_type="min", start_color="E2F0D9", end_type="max", end_color="F4CCCC"),
    )

    label_rows = []
    for variant in VARIANTS:
        if variant not in profile["variants"]:
            continue
        state = profile["variants"][variant]
        total = state["transaction_rows"]
        positives = state["positive_transactions"]
        negatives = state["negative_transactions"]
        label_rows.append(
            [
                variant,
                total,
                positives,
                positives / total,
                negatives / total,
                negatives / total,
                0.0,
                0,
                0,
                state["positive_sender_account_count"],
                state["positive_receiver_account_count"],
                state["positive_endpoint_account_count"],
                "Derived endpoint roll-up: any involvement in a label-1 transaction; not IBM account truth",
            ]
        )
    _table_sheet(
        workbook,
        "Label & Baseline",
        "Class imbalance and an all-negative transaction baseline. Accuracy is misleading here: the no-alert baseline has high accuracy but zero recall.",
        [
            "Variant", "Transactions", "Positive Transactions", "Positive Rate",
            "Negative Rate", "All-Negative Baseline Accuracy", "Baseline Recall",
            "Baseline Precision (no alerts)", "Baseline F1", "Distinct Positive Sender Accounts",
            "Distinct Positive Receiver Accounts", "Distinct Positive Endpoint Accounts",
            "Account Roll-up Definition",
        ],
        label_rows,
        [14, 16, 18, 14, 14, 23, 14, 23, 12, 22, 24, 24, 65],
        "LabelBaselineTable",
    )

    monthly_map: dict[tuple[str, str], dict[int, int]] = defaultdict(lambda: {0: 0, 1: 0})
    for row in profile["aggregates"]["monthly"]:
        monthly_map[(row["variant"], row["month"])][int(row["label"])] += int(row["count"])
    month_rows: list[list[Any]] = []
    for variant, month in sorted(monthly_map):
        counts = monthly_map[(variant, month)]
        total = counts[0] + counts[1]
        within_primary = (
            PRIMARY_START[variant].isoformat()[:7] <= month <= PRIMARY_END[variant].isoformat()[:7]
        )
        month_rows.append(
            [variant, month, total, counts[1], counts[1] / total if total else 0,
             "Primary period month" if within_primary else "Outside stated primary period"]
        )
    _table_sheet(
        workbook,
        "Time Profile",
        "Monthly counts across the actual transaction timestamps. Use temporal splits; specifically inspect positive rows after each variant's stated primary end date.",
        ["Variant", "Month", "Transactions", "Label 1", "Positive Rate", "Period Note"],
        month_rows,
        [15, 16, 18, 15, 16, 34],
        "MonthlyProfileTable",
    )

    format_map = _count_by_variant_label(
        profile["aggregates"]["formats"], "format", "count"
    )
    format_rows = []
    for (variant, payment_format), counts in sorted(format_map.items()):
        total = counts[0] + counts[1]
        format_rows.append(
            [variant, payment_format, total, counts[0], counts[1],
             counts[1] / total if total else 0, counts[1] / max(1, sum(
                 row["count"] for row in profile["aggregates"]["formats"]
                 if row["variant"] == variant and row["label"] == 1
             ))]
        )
    _table_sheet(
        workbook,
        "Payment Formats",
        "Payment Format is a channel/format field, not an account type or risk verdict. All counts are full population.",
        ["Variant", "Payment Format", "Transactions", "Label 0", "Label 1",
         "Positive Rate Within Format", "Share of Variant Label-1 Transactions"],
        format_rows,
        [15, 22, 17, 15, 15, 25, 31],
        "PaymentFormatTable",
    )

    currency_groups: dict[tuple[str, str, str], dict[int, int]] = defaultdict(lambda: {0: 0, 1: 0})
    for row in profile["aggregates"]["currencies"]:
        currency_groups[(row["variant"], row["side"], row["currency"])][int(row["label"])] += int(row["count"])
    currency_rows = []
    for (variant, side, currency), counts in sorted(currency_groups.items()):
        total = counts[0] + counts[1]
        currency_rows.append(
            [variant, side, currency, total, counts[0], counts[1],
             counts[1] / total if total else 0]
        )
    _table_sheet(
        workbook,
        "Currency Profile",
        "Currencies are counted on their own transaction side. Never add or compare raw amounts from different currencies as if they shared a unit.",
        ["Variant", "Amount Side", "Currency", "Transactions", "Label 0", "Label 1", "Positive Rate"],
        currency_rows,
        [15, 16, 23, 18, 16, 16, 16],
        "CurrencyProfileTable",
    )

    amount_rows = []
    for item in sorted(
        profile["aggregates"]["amounts"],
        key=lambda x: (x["variant"], x["side"], x["currency"], x["label"]),
    ):
        count = int(item["count"])
        mean = float(item["sum"]) / count if count else 0
        variance = max(0.0, float(item["sum_squares"]) / count - mean * mean) if count else 0
        amount_rows.append(
            [
                item["variant"], item["side"], item["currency"], item["label"], count,
                float(item["min"]), float(item["max"]), mean, math.sqrt(variance),
                item["approximate_p05"], item["approximate_median"], item["approximate_p95"],
                item["percentile_sample_n"],
            ]
        )
    _table_sheet(
        workbook,
        "Amount Profile",
        "Min/max/mean/population SD use all rows. P05/median/P95 are approximate random samples, separately stratified by variant, side, currency and label. A currency is never converted or pooled with another.",
        ["Variant", "Side", "Currency", "IBM Label", "Rows", "Exact Min",
         "Exact Max", "Exact Mean", "Approx. Population SD", "Approx. Sample P05",
         "Approx. Sample Median", "Approx. Sample P95", "Sample N"],
        amount_rows,
        [15, 14, 22, 12, 16, 17, 17, 18, 22, 20, 22, 20, 14],
        "AmountProfileTable",
    )

    pair_rows = []
    for row in sorted(
        profile["aggregates"]["currency_pairs"],
        key=lambda x: (x["variant"], x["payment_currency"], x["receiving_currency"], x["label"]),
    ):
        pair_rows.append(
            [row["variant"], row["payment_currency"], row["receiving_currency"],
             row["label"], row["count"]]
        )
    _table_sheet(
        workbook,
        "Currency Pairs",
        "Exact frequency of payment-to-receiving currency combinations, split by IBM label. This is a pair distribution, not an FX conversion.",
        ["Variant", "Payment Currency", "Receiving Currency", "IBM Label", "Transactions"],
        pair_rows,
        [15, 23, 24, 13, 18],
        "CurrencyPairTable",
    )

    account_rows = []
    bank_rows = []
    for variant in VARIANTS:
        if variant not in profile["variants"]:
            continue
        state = profile["variants"][variant]
        account_rows.append(
            [
                variant, state["account_rows"], state["unique_bank_id_source_values"],
                state["unique_bank_names"],
                state["approximate_unique_entity_ids"],
                state["account_rows"] / max(1, state["unique_bank_id_source_values"]),
                "Not supplied",
                "Bank Name, Bank ID, Account Number, Entity ID, Entity Name",
            ]
        )
        for bank in state["bank_account_counts"]:
            bank_rows.append(
                [variant, bank["bank_id"], bank["bank_name"], bank["account_rows"]]
            )
    _table_sheet(
        workbook,
        "Account Profile",
        "Account dimensions do not contain account product/type or account currency; see Account Type Guide for definitions and evidence requirements. Unique entity counts are HLL estimates (~1.6% typical relative error at this configuration); account rows and banks are exact.",
        ["Variant", "Account Dimension Rows", "Bank ID Values", "Bank Names",
         "Approx. Unique Entity IDs", "Accounts per Bank ID (avg.)",
         "Account Type", "Available Source Account Fields"],
        account_rows,
        [15, 22, 17, 15, 25, 25, 20, 70],
        "AccountProfileTable",
    )
    account_type_rows = [
        ["Personal/retail deposit", "Product + owner", "Deposit account held for an individual for storing funds and making payments, such as current/checking or savings.", "Effective-dated bank product code and verified customer type.", "Not supplied; Entity Name is not a verified customer type.", "Do not infer from payment format, names, or amounts."],
        ["Business operating deposit", "Product + owner", "Deposit account used by a company or organization for operating receipts and payments.", "Account product plus verified legal-entity/customer master.", "Not supplied.", "Entity ID/Name alone does not provide a verified legal classification."],
        ["Credit/loan account", "Credit product", "Account or facility recording borrowing, card credit, repayments, or a credit obligation.", "Lending/card product master and account relationship.", "Not supplied.", "Payment Format is not an account product code."],
        ["Wallet/e-money/prepaid", "Product", "Stored-value or electronic-money product used to hold or move value; regulatory treatment varies.", "Provider/product registry and effective-dated account/wallet master.", "Not supplied.", "Do not infer from a transaction channel or counterparty name."],
        ["Clearing account", "Operational ledger role", "Temporary ledger account collecting, routing, or reconciling postings while a payment/accounting process completes; may be internal GL, not a customer deposit account.", "GL/account-purpose code, owning process, reconciliation rules, and effective dates.", "Not supplied; no account-purpose, GL, or balance fields.", "High volume, many counterparties, and rapid pass-through are review clues only."],
        ["Settlement account", "Operational/payment role", "Account used by payment participants/institutions to discharge payment obligations, often following clearing or netting.", "Settlement-system participation/role and settlement documentation.", "Not supplied.", "No settlement-system or account-purpose field is present."],
        ["Correspondent (Nostro/Vostro)", "Inter-institutional relationship", "An account/service relationship where one financial institution serves another. Nostro and Vostro describe the same relationship from opposite parties' viewpoints.", "Verified FI identities, account relationship, and correspondent agreement.", "Not supplied; bank identifiers are synthetic.", "Do not infer real FIs or cross-border correspondent roles from IBM IDs."],
        ["Suspense account", "Operational ledger role", "Temporary posting location for items whose final account, purpose, or allocation is unresolved, pending investigation/reconciliation.", "GL/account-purpose code and unresolved-item/reconciliation records.", "Not supplied; no posting status or balance data.", "Temporary or unusual movement alone does not prove a suspense account."],
        ["Transit/intermediary account", "Operational role", "Account or ledger used temporarily to route value between processing steps or destinations.", "Account/GL metadata and documented processing flow.", "Not supplied.", "A graph path or pass-through pattern alone cannot establish the role."],
        ["Escrow/trust/client-money", "Legal/holding role", "Funds held by a third party or fiduciary for another party or purpose under defined terms; legal treatment varies by jurisdiction.", "Account designation, legal agreement, beneficial-party records, and local rules.", "Not supplied.", "Owner or transaction names do not establish fiduciary status."],
        ["Merchant/acquirer settlement", "Payment role", "Merchant proceeds collected or paid out through an acquiring/payment provider, often in batches after fees or adjustments.", "Merchant/acquirer identifiers, payment-system role, and settlement records.", "Not supplied; no merchant/acquirer field.", "Payment Format does not identify the merchant or acquirer."],
        ["Omnibus/pooled (including FBO)", "Ownership/ledger structure", "One external account holds pooled funds for multiple underlying customers whose balances are tracked in a separate subledger.", "Account title, intermediary/custodian records, underlying-customer subledger, and reconciliations.", "Not supplied; no subledger or beneficial-owner relationship.", "Do not infer pooled ownership from account activity alone."],
    ]
    _table_sheet(
        workbook,
        "Account Type Guide",
        "Definitions are practical, not universal legal classifications. Product type, owner type, and operational role can overlap. For IBM, all these classifications remain unknown because source account-purpose/GL fields are absent.",
        ["Term", "Dimension", "Practical Definition", "Evidence Needed", "IBM Source Status", "Interpretation Control"],
        account_type_rows,
        [28, 28, 76, 70, 58, 66],
        "AccountTypeGuideTable",
    )
    bank_identifier_rows = [
        ["IBM Bank ID", "Synthetic benchmark-local bank key", "Join to the source Bank Name in the account dimension, scoped to dataset variant.", "Dataset Variant + normalized Bank ID + Account Number.", "Not a real national, SWIFT/BIC, or global identifier; never external-match the ID."],
        ["IBM Bank Name", "Source-provided display name", "Display the matching account-dimension value as supplied by IBM.", "Join transaction endpoint to account dimension using variant, bank ID, and account number.", "Synthetic labels may resemble real names or mention countries; this does not verify a real institution."],
        ["BIC (SWIFT)", "International business-identifier scheme", "Identifies a business party and, where present, branch/location; common forms use 8 or 11 characters.", "Authoritative BIC directory, source jurisdiction, effective date, and verification.", "Not an account number, bank-account owner, or account-purpose flag. Official directory access may be licensed."],
        ["LEI", "Legal entity identifier", "20-character identifier for a legal entity; GLEIF records can include BIC mappings for covered entities.", "GLEIF API/Golden Copy and legal-name/jurisdiction match.", "Not a payment-routing code or account ID; not every entity has an LEI."],
        ["National routing code", "Country/scheme-specific", "Domestic institution, branch, or payment-routing code (for example ABA, UK sort code, or South African universal/branch code).", "Official national regulator or payment-system directory for the relevant country and scheme.", "No universal format; keep scheme and jurisdiction with the value."],
        ["IBAN", "Jurisdiction-specific account-number standard", "Structured account number used by participating jurisdictions, with country-specific components and check digits.", "Validate against the relevant country format and institution source.", "Not a global bank directory or a substitute for validating an institution/account."],
        ["Recommended enrichment crosswalk", "Governed external reference", "Store raw code, code scheme, jurisdiction, source/matched bank name, BIC/LEI where available, branch, valid-from/to, source/version, retrieval date, match method, and review status.", "Use trusted, effective-dated source references and human review for ambiguous matches or mergers.", "Do not fuzzy-match synthetic IBM Bank IDs to real-world codes."],
    ]
    _table_sheet(
        workbook,
        "Bank Identifier Guide",
        "For this benchmark, resolve names with IBM's account dimension. International and domestic bank identifiers describe different entities/schemes; there is no universal code that can map these synthetic Bank IDs to real banks.",
        ["Identifier", "Scope", "Meaning", "Recommended Resolution", "Limits / IBM Application"],
        bank_identifier_rows,
        [29, 36, 72, 73, 78],
        "BankIdentifierGuideTable",
    )
    _table_sheet(
        workbook,
        "Bank Profile",
        "Exact account-dimension record counts by source Bank ID and Bank Name. These are benchmark-local identifiers/names, not verified real banks. Normalize bank IDs only for a documented account-dimension join while preserving original text.",
        ["Variant", "Source Bank ID", "Bank Name", "Account Rows"],
        bank_rows,
        [15, 19, 34, 17],
        "BankProfileTable",
    )

    pattern_rows = []
    for variant, families in profile["pattern_annotations"].items():
        for family in sorted(families):
            item = families[family]
            pattern_rows.append(
                [variant, family, item["attempts"], item["annotation_lines"],
                 "Annotation lines are not guaranteed unique transaction IDs; "
                 "one row may occur in multiple attempts. Not all label-1 rows are in these patterns."]
            )
    _table_sheet(
        workbook,
        "Pattern Catalogue",
        "Counts are derived from each retained IBM pattern text file. Interpret the names as simulator scenario shapes, not complete real-world typologies.",
        ["Variant", "IBM Scenario Family", "Attempt Blocks", "Annotated Transaction Lines", "Interpretation Note"],
        pattern_rows,
        [15, 25, 17, 28, 90],
        "PatternCatalogueTable",
    )

    field_rows = [
        ("Timestamp", "Transaction", "Source", "Transaction DateTime", "Direct; source has minute precision and no timezone column.", "Confirm source timezone and target batch cutoff before production use."),
        ("From Bank", "Transaction", "Source", "Account Key (sender side); Opposite Bank (receiver side)", "Bank ID text; leading zeroes can differ from account dimension formatting.", "Keep raw code; normalized numeric key is only a join helper."),
        ("Account", "Transaction", "Source", "Account Key (sender side); Opposite Account Number (receiver side)", "Sender account number.", "Join with variant and bank ID; number alone is not unique."),
        ("To Bank", "Transaction", "Source", "Opposite Bank (sender side); Account Key (receiver side)", "Receiver bank ID.", "Preserve text; no country code provided."),
        ("Account.1", "Transaction", "Source", "Opposite Account Number (sender side); Account Key (receiver side)", "Receiver account number; source CSV repeats the Account header.", "Use receiver role explicitly."),
        ("Amount Paid", "Transaction", "Source", "Transaction Amount in Original Currency (sender event)", "Sender-side debit in Payment Currency.", "Correct side and currency mapping."),
        ("Payment Currency", "Transaction", "Source", "Original Currency Code (sender event)", "Currency of Amount Paid.", "Not the sender account's configured currency."),
        ("Amount Received", "Transaction", "Source", "Transaction Amount in Original Currency (receiver event)", "Receiver-side credit in Receiving Currency.", "Use a receiver event row to map this side."),
        ("Receiving Currency", "Transaction", "Source", "Original Currency Code (receiver event)", "Currency of Amount Received.", "Not the receiver account's configured currency."),
        ("Payment Format", "Transaction", "Source", "Transaction Type (proxy only)", "Recorded channel/format.", "Not necessarily the target system's business transaction type."),
        ("Is Laundering", "Transaction", "Benchmark label", "Ground-truth file only", "Synthetic label for retrospective evaluation.", "Never expose in scoring input/features."),
        ("Dataset Variant", "Staged", "Derived", "Dataset Variant (benchmark only)", "Variant provenance.", "Keep for benchmark lineage; omit from real-client inference input."),
        ("Transaction Key", "Staged", "Generated surrogate", "Unique Identifier for the transaction", "Variant + one-based source row key, generated by staging.", "Not an IBM-issued transaction identifier."),
        ("Bank Name", "Account", "Source account dimension", "Optional account/bank lookup", "Bank Name from account file.", "Join via variant, normalized Bank ID, Account Number."),
        ("Bank ID", "Account", "Source account dimension", "Account Key", "Account's bank identifier.", "Raw text retained; normalize only for composite matching."),
        ("Account Number", "Account", "Source account dimension", "Account Key", "Account identity within a bank and variant.", "Account number alone is not globally unique."),
        ("Entity ID", "Account", "Source account dimension", "Customer/entity key (enrichment)", "Source entity/party identifier.", "Not an account-type attribute."),
        ("Entity Name", "Account", "Source account dimension", "Customer/entity name (enrichment)", "Source-provided party name.", "Synthetic identity; no sanctions list matching validation."),
        ("Account Type", "Frame target", "Not provided", "Account profile template", "No checking/savings/business/customer-account category exists in source schema.", "Obtain from Frame/client account master; leave blank in IBM fixture."),
        ("Account Currency", "Frame target", "Not provided", "Account profile template", "No account-level currency is supplied.", "Do not assume ZAR; calculate account-currency amount only with approved enrichment/rule."),
        ("Transaction Amount in Account Currency", "Frame target", "Not provided", "Transaction template", "Cannot be calculated from IBM fields alone.", "Require account currency and effective FX rate/source."),
        ("Transaction Branch", "Frame target", "Not provided", "Transaction/account template", "No branch or channel location attribute.", "Obtain from the target bank system."),
        ("Opposite Country Code", "Frame target", "Not provided", "Transaction template", "No transaction-side country fields.", "Do not infer country from bank ID or a synthetic name."),
        ("Transaction Direction", "Frame target", "Derived role", "Two account-event rows per transaction", "Sender = Debit; receiver = Credit for account-centric event representation.", "This is a target modelling transformation; original transaction is one directed transfer."),
    ]
    _table_sheet(
        workbook,
        "Field Dictionary",
        "Every proposed mapping is explicit. “Not provided” means the feature cannot be truthfully generated from IBM alone.",
        ["Field", "Record Area", "Source Status", "FrameAML Mapping", "Meaning / Transformation", "Control / Limitation"],
        [field_rows[i:i+1][0] for i in range(len(field_rows))],
        [25, 19, 25, 47, 64, 75],
        "FieldDictionaryTable",
    )

    _table_sheet(
        workbook,
        "CSV Mapping",
        "Account-centric target mapping aligned to the supplied Frame attribute sketch. The IBM transaction is one directed row; create two account-perspective rows only when the receiving system requires one row per account-side event.",
        ["Frame Target Attribute", "IBM Source", "Availability", "Transformation / Example", "Required Enrichment / Caveat"],
        [
            ["Account Key", "Dataset Variant + Bank ID + Account Number", "Derivable", "variant|normalized-bank-id|account-number. Separate keys for sender and receiver.", "Account number alone is not unique; preserve source text in separate fields."],
            ["Transaction Amount in Account Currency", "None", "Unavailable", "Keep blank in the IBM fixture.", "Needs a real account-currency master and approved FX conversion. Never assume ZAR."],
            ["Transaction Amount in Original Currency", "Amount Paid or Amount Received", "Available by role", "Sender event uses Amount Paid; receiver event uses Amount Received.", "Do not mix the debit and credit sides."],
            ["Original Currency Code", "Payment Currency or Receiving Currency", "Available by role", "Use Payment Currency for sender; Receiving Currency for receiver.", "Currencies can differ within one transaction."],
            ["Transaction Branch", "None", "Unavailable", "Blank.", "Obtain from bank transaction/account platform."],
            ["Transaction DateTime", "Timestamp", "Available", "Preserve source timestamp.", "No timezone field; define target timezone semantics."],
            ["Opposite account number", "Account.1 or Account", "Available by role", "Sender event opposite=receiver; receiver event opposite=sender.", "Keep same transaction unique ID on both rows."],
            ["Opposite bank", "To Bank or From Bank", "Available by role", "Sender event opposite bank=To Bank; receiver event=From Bank.", "Identifiers are synthetic bank codes, not country codes."],
            ["Opposite country code", "None", "Unavailable", "Blank.", "Must be supplied from trusted reference data."],
            ["Transaction Type", "Payment Format", "Proxy only", "Pass through channel string (e.g., ACH) in development.", "Do not represent Payment Format as a full business transaction type."],
            ["Transaction Direction", "Sender/receiver role", "Derived", "Sender=Debit; receiver=Credit.", "Only for account-event schema; one-way fact schema should keep direction as sender→receiver."],
            ["Unique Identifier for the transaction", "Transaction Key", "Generated trace key", "Dataset Variant + source row ordinal.", "Not IBM-issued; generate new immutable ID for real source transactions."],
            ["Account Type", "None", "Unavailable", "Not in transaction or account source schemas.", "Add from the Frame/client account master; retain effective dates."],
            ["IBM Benchmark Label", "Is Laundering", "Evaluation only", "Separate ground-truth file by transaction ID.", "Never include in model or operational scoring features."],
        ],
        [37, 34, 20, 66, 68],
        "FrameCsvMappingTable",
    )

    protocols = [
        ("Purpose & unit", "Declare either transaction classification (one directed transfer per ID) or account-side event testing (two rows per transfer). Report metrics for the chosen unit; do not silently mix them."),
        ("Data partitions", "Use HI-Small for a first full-population baseline; compare HI and LI separately. Variants are independent, not nested. Reserve chronological train/validation/test partitions; record cutoffs."),
        ("Temporal leakage", "Kaggle's description says post-period transactions are laundering, but this full scan finds 56,691 rows after the stated cutoffs and only 33,583 label-1 rows. Reconcile the boundary, data version and cited discussion before using temporal splits; never treat date alone as ground truth."),
        ("Baseline first", "Measure a transparent rules baseline before graph/ML. Include all-negative/no-alert baseline to expose the accuracy trap. Then compare rules, graph features, ML, and fused scores under the same scope/split."),
        ("Metrics", "Publish precision, recall, minority-class F1, PR-AUC, false positives, alert volume, processing time and retained-positive recall. Add account-level network metrics and clearly label their separate evaluation unit."),
        ("Pattern tests", "For fan-in/out, cycles, stack and gather/scatter, emit triggering transaction IDs, ordered hops, time gaps, involved accounts, per-currency amounts and the rule version. Preserve alternative benign explanations."),
        ("Avoid leakage", "Remove Is Laundering, pattern-file labels, post-outcome investigation fields and direct IDs from model features. IDs may be used for grouping/graph topology, but not as memorization features."),
        ("Deploy boundary", "This benchmark cannot validate sanctions name matching, TF-specific remittance/NPO typologies, proliferation-financing trade controls or real client alert quality. Follow with institution-specific historical data and analyst review."),
        ("Thresholds", "Treat candidate time windows and thresholds as experiments. Calibrate by currency, account/entity baseline, alert capacity and cost of misses vs false alerts; do not claim regulatory thresholds from this synthetic dataset."),
    ]
    _table_sheet(
        workbook,
        "Test Protocol",
        "Recommended preparation for a defensible FrameAML second test suite.",
        ["Protocol Area", "Required Test Design"],
        [[name, instruction] for name, instruction in protocols],
        [28, 120],
        "TestProtocolTable",
    )

    source_rows = [
        [
            "Kaggle dataset page",
            "Dataset description and downloadable files",
            "Six independent variants (HI/LI × Small/Medium/Large); HI is higher-illicit-rate and LI lower; chronological 60/20/20 split is suggested; stated date ranges are primary periods; eight named pattern families are listed, but not all positive transactions follow them. The live file-list API shows 18 files: transactions, accounts and patterns for each variant, although the page body lists only the 12 transaction/pattern files. Kaggle's description says post-period transactions are laundering; this full scan counts 56,691 post-cutoff rows but only 33,583 label-1 rows, so reconcile the cutoff/version/discussion before a temporal split.",
            "https://www.kaggle.com/datasets/ealtman2019/ibm-transactions-for-anti-money-laundering-aml",
            "Fetched live; Kaggle API metadata reports Community Data License Agreement - Sharing 1.0 and ~41.6 GB total dataset size.",
        ],
        [
            "Kaggle dataset file-list API",
            "Current downloadable-file inventory",
            "Confirms 18 files: one transaction CSV, one account CSV and one pattern text file for each of HI/LI Small, Medium and Large. This is more complete than the page body's older list of 12 files.",
            "https://www.kaggle.com/api/v1/datasets/list/ealtman2019/ibm-transactions-for-anti-money-laundering-aml",
            "Fetched live; returned per-file names, byte sizes and creation timestamps.",
        ],
        [
            "IBM AMLworld paper (NeurIPS 2023)",
            "Generator design, limits and intended benchmark use",
            "Synthetic multi-agent financial ecosystem; models placement, layering and integration; complete simulated ground truth supports benchmarking; transactions naturally form a directed multigraph; synthetic results do not establish real-bank performance.",
            "https://arxiv.org/abs/2306.16424",
            "Fetched live at arXiv HTML; paper describes generator calibration and provides dataset context.",
        ],
        [
            "IBM Multi-GNN repository",
            "Reference AML graph-model code",
            "Implements GIN, GAT, PNA and RGCN variants plus edge MLPs, reverse message passing, ego IDs and port numbering; repository states experiments use Kaggle transaction files and the AMLworld and directed-multigraph papers.",
            "https://github.com/IBM/Multi-GNN",
            "README fetched live; use as a reproducibility reference, not a production-ready detector or proof of client performance.",
        ],
        [
            "Directed multigraph GNN paper (AAAI 2024)",
            "Directed transaction-network modeling reference",
            "Studies adaptations for directed multigraphs and reports improved minority-class F1 in its evaluated tasks; results are paper-specific and should be re-run under FrameAML's own split, unit and feature controls.",
            "https://arxiv.org/abs/2306.11586",
            "Abstract and citation metadata fetched live.",
        ],
        [
            "IBM Box legacy dataset page",
            "Earlier generator release reference",
            "The Kaggle description says the generator was significantly improved since this earlier release. Do not mix legacy data with current Kaggle variants without versioning and schema checks.",
            "https://ibm.box.com/v/AML-Anti-Money-Laundering-Data",
            "Live page redirected to Box but the document content was not accessible through public fetch; Kaggle page confirms its reference and caveat.",
        ],
        [
            "IBM Box PDF file link",
            "Legacy AML dataset PDF",
            "Direct PDF file URL supplied by the user.",
            "https://ibm.ent.box.com/v/AML-Anti-Money-Laundering-Data/file/780515045707",
            "Box returned its generic login/application page rather than the PDF content; not represented as reviewed.",
        ],
        [
            "Kaggle discussion 427517",
            "Post-primary-period positive-label context",
            "The dataset page points to a discussion response for the post-period label issue. This full scan finds 56,691 post-cutoff rows and 33,583 positives, which does not reproduce an all-positive post-period result under the date boundary used here.",
            "https://www.kaggle.com/datasets/ealtman2019/ibm-transactions-for-anti-money-laundering-aml/discussion/427517",
            "Direct page request returned 404; the Kaggle dataset description refers readers to this discussion. Reconcile the discrepancy before using the date boundary for evaluation.",
        ],
        [
            "IBM AMLSim repository",
            "Pattern-family background",
            "IBM AMLSim is a synthetic transaction-graph simulator; Kaggle states its eight pattern families were introduced by this simulator.",
            "https://github.com/IBM/AMLSim",
            "Repository README fetched live; this is background, not the current Kaggle data generator.",
        ],
        [
            "Related Kaggle credit-card transactions",
            "Separate synthetic fraud-test data",
            "IBM describes more than 20 million purchase transactions from a simulated world covering 2,000 synthetic US consumers, decades of purchases, and multiple cards per consumer. Kaggle's file API lists credit_card_transactions-ibm_v2.csv (~2.35 GB), User0_credit_card_transactions.csv, sd254_cards.csv, and sd254_users.csv. The page describes fraud detection and purchase-pattern analysis; this is not an AML laundering or inter-account transfer label set and cannot validate AML typologies. Consider it only as a separately governed fraud-test suite.",
            "https://www.kaggle.com/datasets/ealtman2019/credit-card-transactions",
            "Page and Kaggle file-list API fetched live. The page's description includes Apache License 2.0 terms; Kaggle metadata says 'Other (specified in description)'. Verify the applicable license terms before redistribution. Dataset is not included in the IBM AML population scan.",
        ],
        [
            "Kaggle AML baseline notebook",
            "Community analysis reference",
            "The current page identifies 'Baseline - IBM Transactions for AML' by Gaurav Bhattacharya, published/modified 17 May 2024, and links it to this AML dataset. The notebook cell code and execution outputs were not retrievable from the public page/API used here, so no algorithm, metric, or preprocessing claim is attributed to it. Review its actual code and data version before adopting any steps.",
            "https://www.kaggle.com/code/gbiamgaurav/baseline-ibm-transactions-for-aml",
            "Public metadata fetched; code page renders a JavaScript viewer. Direct kernel-code/scriptcontent endpoints returned 404 without a retrievable notebook payload.",
        ],
        [
            "GLEIF API",
            "Public legal-entity and identifier cross-reference",
            "GLEIF exposes LEI record search and BIC mappings for covered legal entities. A live query for NEDSZAJJXXX returned an active Nedbank Limited LEI record; this is an example of resolving a real BIC, not mapping IBM synthetic Bank IDs.",
            "https://api.gleif.org/api/v1/lei-records?filter%5Bbic%5D=NEDSZAJJXXX",
            "Live API response fetched 26 Oct 2026. Coverage is limited to LEI records/relationships in GLEIF data; not a universal routing directory.",
        ],
        [
            "SWIFT BIC standard",
            "International business-identifier scheme reference",
            "BIC is a distinct identifier scheme for business parties/locations; it is not an account number or account-type flag. Use an authoritative BIC source for any real-data crosswalk.",
            "https://www.swift.com/standards/data-standards/bic-business-identifier-code",
            "Official site returned HTTP 403 to automated retrieval; no SWIFT directory records were downloaded or matched to IBM.",
        ],
    ]
    _table_sheet(
        workbook,
        "Research Sources",
        "Online source pages were checked during preparation. Access limits are disclosed rather than inferred; links are references, not endorsements.",
        ["Source", "Why It Matters", "Verified Takeaways", "URL", "Access / Scope Note"],
        source_rows,
        [31, 35, 100, 84, 85],
        "ResearchSourcesTable",
    )

    quality_rows = []
    for variant in VARIANTS:
        if variant not in profile["variants"]:
            continue
        state = profile["variants"][variant]
        for source_name, observed, expected in (
            ("Transaction", state["transaction_null_counts"], state["manifest_transaction_null_counts"]),
            ("Account", state["account_null_counts"], state["manifest_account_null_counts"]),
        ):
            for column in sorted(set(observed) | set(expected)):
                quality_rows.append(
                    [
                        variant, source_name, column, int(observed.get(column, 0)),
                        int(expected.get(column, 0)),
                        "PASS" if int(observed.get(column, 0)) == int(expected.get(column, 0)) == 0 else "REVIEW",
                    ]
                )
        quality_rows.extend(
            [
                [variant, "Transaction", "Scanned rows = manifest", state["transaction_rows"],
                 state["expected_transaction_rows"], "PASS" if state["transaction_rows"] == state["expected_transaction_rows"] else "FAIL"],
                [variant, "Account", "Scanned rows = manifest", state["account_rows"],
                 state["expected_account_rows"], "PASS" if state["account_rows"] == state["expected_account_rows"] else "FAIL"],
                [variant, "Account", "Composite key uniqueness", "Unique in staging quality report",
                 "Dataset Variant + numeric-normalized Bank ID + Account Number", "PASS" if state["account_composite_key_unique_per_staging_validation"] else "FAIL"],
                [variant, "Join coverage", "Transaction endpoint → account dimension", "Not measured in this population scan",
                 "Measure sender and receiver match rates before name-level use", "ACTION REQUIRED"],
            ]
        )
    _table_sheet(
        workbook,
        "Data Quality",
        "Observed nulls and row-counts are checked against existing quality reports. Account-key uniqueness is from staging validation; transaction-to-account join coverage remains an explicit next check.",
        ["Variant", "Table", "Check", "Observed", "Expected / Rule", "Status"],
        quality_rows,
        [15, 17, 42, 32, 75, 18],
        "DataQualityTable",
    )

    ryno_source_hash = _sha256_file(RYNO_TRANSACTION_PATH)
    _table_sheet(
        workbook,
        "Ryno Comparison",
        "Descriptive comparison of the complete HI-Small source with Ryno's top-20-bank UDM output and single-bank sanitised output. Payments and account-side legs are different units; the transformed subsets are not comparable benchmark populations.",
        ["Scope", "Measure", "Value", "Unit", "Interpretation / Limitation"],
        _ryno_comparison_rows(profile, ryno_source_hash),
        [34, 37, 34, 22, 100],
        "RynoComparisonTable",
    )
    _table_sheet(
        workbook,
        "Ryno Field Mapping",
        "All 12 attributes from the supplied sanitised transaction specification, in its requested order. Direct source fields are separated from generated keys, leg transformations, and unverified assumptions; validate names against the target NICE Actimize UDM release.",
        [
            "Required Attribute",
            "JSON Key",
            "HI-Small Source",
            "Ryno Transformation",
            "Reliability / Action",
        ],
        [
            [
                "Unique Identifier for the transaction",
                "transactionId",
                "No source transaction ID; source row position is available.",
                "Deterministic `HS-<row id>-<D|C>` ID generated per account leg.",
                "Generated surrogate, not IBM-issued. Retain the source row ID for traceability.",
            ],
            [
                "Account Key",
                "accountKey",
                "In-scope Bank ID and account number (`Account` or `Account.1`).",
                "Composite `<bankId>|<accountNumber>` for the account on this leg.",
                "Bank ID plus account number is required; account number alone is not unique. Bank IDs are synthetic benchmark keys.",
            ],
            [
                "Transaction Date/Time",
                "transactionDateTime",
                "Timestamp, minute precision, no timezone field.",
                "Formatted as `YYYY-MM-DDTHH:MM:SS` and treated as local time.",
                "Timezone semantics are unknown. Confirm the source/business timezone before production use.",
            ],
            [
                "Transaction Type",
                "transactionType",
                "Payment Format plus debit/credit perspective.",
                "Mapped from payment format × direction (for example ACH_OUT or CASH_DEP).",
                "Derived transaction code, not a source business transaction code. Confirm the mapping against the target system.",
            ],
            [
                "Transaction Direction",
                "transactionDirection",
                "From/To bank and account roles.",
                "One D debit leg for the payer and/or one C credit leg for the beneficiary.",
                "Changes the unit from payment row to account event. Linked legs must retain a common source reference.",
            ],
            [
                "Transaction Amount in Original Currency",
                "transactionAmountOriginal",
                "D: Amount Paid. C: Amount Received.",
                "Selects the amount belonging to the account-side leg.",
                "Direct value by side. Do not mix payment and receiving sides or currencies.",
            ],
            [
                "Original Currency Code",
                "originalCurrencyCode",
                "D: Payment Currency. C: Receiving Currency.",
                "Mapped to ISO-style codes; Bitcoin is represented as XBT.",
                "Direct currency by side with code normalization. Not account currency.",
            ],
            [
                "Transaction Amount in Account Currency",
                "transactionAmountAccountCcy",
                "No account-level currency or account-currency amount is supplied.",
                "Treats every account as ZAR and applies implied HI-Small FX × approximate 17.50 ZAR/USD.",
                "Synthetic estimate; 17.50 is not an official historical rate. Confirm account currency and approved FX source.",
            ],
            [
                "Transaction Branch",
                "transactionBranch",
                "No branch field is supplied.",
                "Fabricated deterministic home branch per account; teller formats use home branch and electronic formats use virtual B000 by default.",
                "Not a source fact. Replace with trusted branch mapping; test the chosen branch-mode rule.",
            ],
            [
                "Opposite account number",
                "oppositeAccountNumber",
                "The opposite endpoint's `Account` or `Account.1`.",
                "Copies the other party's account number; self-transfer points to the same account.",
                "Direct counterparty value, but only one side's leg is emitted for an external account.",
            ],
            [
                "Opposite bank",
                "oppositeBank",
                "The opposite endpoint's `To Bank` or `From Bank`.",
                "Bank ID is normalized by removing zero padding by default; optional bank-name output.",
                "Synthetic benchmark-local bank ID, not a BIC, LEI, or national routing code.",
            ],
            [
                "Opposite country code",
                "oppositeCountryCode",
                "No transaction-side country code is supplied.",
                "Inferred from account dimension bank name with name-based country rules.",
                "Derived heuristic, not source country. Confirm with trusted institution reference data.",
            ],
        ],
        [38, 29, 48, 77, 80],
        "RynoFieldMappingTable",
    )

    guide_title_row = len(start_values) + 5
    start.merge_cells(f"A{guide_title_row}:B{guide_title_row}")
    guide_title = start.cell(guide_title_row, 1, "Workbook sheet guide")
    guide_title.fill = PatternFill("solid", fgColor=NAVY)
    guide_title.font = Font(name="Aptos Display", size=13, bold=True, color=WHITE)
    guide_title.alignment = Alignment(vertical="center")
    start.row_dimensions[guide_title_row].height = 26
    guide_header_row = guide_title_row + 1
    for column, value in enumerate(("Sheet", "What it contains / how to interpret it"), 1):
        cell = start.cell(guide_header_row, column, value)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN_GREY)
    start.row_dimensions[guide_header_row].height = 30

    guide_rows: list[tuple[str, str]] = []
    for sheet in workbook.worksheets:
        if sheet.title == "Start Here":
            description = (
                "Executive orientation, population scope, key limitations, reading order, "
                "and this sheet-by-sheet index."
            )
        else:
            description = str(sheet["A2"].value or "").strip()
        guide_rows.append((sheet.title, description))
    if include_hismall_raw:
        for sheet_number, (first_row, last_row) in enumerate(
            _raw_sheet_ranges(hismall_rows), 1
        ):
            guide_rows.append(
                (
                    f"HI-Small Raw {sheet_number}",
                    f"Exact original transaction rows {first_row:,}–{last_row:,}, "
                    "with all 11 source columns in original order. Values are text; "
                    "use these tabs to verify individual source payments.",
                )
            )

    for row_number, (sheet_name, description) in enumerate(
        guide_rows, guide_header_row + 1
    ):
        name_cell = start.cell(row_number, 1, sheet_name)
        name_cell.font = Font(name="Aptos", size=9, bold=True, color=NAVY)
        name_cell.fill = PatternFill("solid", fgColor=BLUE)
        description_cell = start.cell(row_number, 2, description)
        description_cell.fill = PatternFill(
            "solid", fgColor="F7F9FC" if row_number % 2 else WHITE
        )
        for cell in (name_cell, description_cell):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN_GREY)
        start.row_dimensions[row_number].height = 46

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    raw_path = output_path.with_name(f".{output_path.stem}.raw.tmp.xlsx")
    compatible_path = output_path.with_name(
        f".{output_path.stem}.compatible.tmp.xlsx"
    )
    try:
        workbook.save(raw_path)
        _rewrite_excel_compatible(
            raw_path,
            compatible_path,
            include_hismall_raw,
            hismall_rows,
            workbook,
        )
        with ZipFile(compatible_path) as package:
            if package.testzip() is not None:
                raise ValueError("The generated workbook package failed its CRC check.")
        check = load_workbook(compatible_path, read_only=True, data_only=True)
        try:
            required_sheets = {
                "Start Here",
                "CSV Mapping",
                "Account Type Guide",
                "Bank Identifier Guide",
                "Ryno Comparison",
                "Ryno Field Mapping",
            }
            if include_hismall_raw:
                required_sheets.update(
                    f"HI-Small Raw {number}"
                    for number in range(1, len(_raw_sheet_ranges(hismall_rows)) + 1)
                )
            if not required_sheets.issubset(check.sheetnames):
                raise ValueError("The generated workbook is missing required sheets.")
            written_totals = list(
                check["Variant Overview"].iter_rows(
                    min_row=5, values_only=True
                )
            )
            written_totals = [row for row in written_totals if row and row[0] is not None]
            if len(written_totals) != len(variant_rows):
                raise ValueError("Variant overview did not persist expected data rows.")
            if (
                sum(int(row[1]) for row in written_totals)
                != profile["staging_totals"]["transaction_rows"]
                or sum(int(row[3]) for row in written_totals)
                != profile["staging_totals"]["positive_transactions"]
            ):
                raise ValueError("Workbook totals do not match the full profile.")
            if include_hismall_raw:
                raw_rows_written = 0
                for sheet_number, (first_row, last_row) in enumerate(
                    _raw_sheet_ranges(hismall_rows), 1
                ):
                    raw_sheet = check[f"HI-Small Raw {sheet_number}"]
                    expected_sheet_rows = last_row - first_row + 1
                    actual_headers = next(
                        raw_sheet.iter_rows(
                            min_row=RAW_SHEET_HEADER_ROW,
                            max_row=RAW_SHEET_HEADER_ROW,
                            min_col=1,
                            max_col=len(HI_SMALL_SOURCE_HEADERS),
                            values_only=True,
                        )
                    )
                    if list(actual_headers) != HI_SMALL_SOURCE_HEADERS:
                        raise ValueError(
                            f"HI-Small Raw {sheet_number} does not preserve the source headers."
                        )
                    raw_rows_written += expected_sheet_rows
                if raw_rows_written != hismall_rows:
                    raise ValueError(
                        f"HI-Small raw tabs contain {raw_rows_written:,} rows; "
                        f"expected {hismall_rows:,}."
                    )
        finally:
            check.close()
        os.replace(compatible_path, output_path)
    finally:
        raw_path.unlink(missing_ok=True)
        compatible_path.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Profile all selected IBM AML Parquet variants and build a FrameAML workbook."
    )
    parser.add_argument(
        "--variants",
        nargs="+",
        choices=VARIANTS,
        default=VARIANTS,
        help="Variants to profile; defaults to all six.",
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=200_000,
        help="Number of Parquet rows scanned per batch.",
    )
    parser.add_argument(
        "--reuse-profile",
        action="store_true",
        help="Rebuild the workbook and CSV outputs from the existing profile JSON without rescanning Parquet.",
    )
    parser.add_argument(
        "--without-hi-small-raw",
        action="store_true",
        help="Build a smaller summary workbook without embedding the five full HI-Small raw-data tabs.",
    )
    args = parser.parse_args()
    if "HI-Small" not in args.variants:
        parser.error("The HI-Small variant is required for the Ryno comparison.")
    if args.reuse_profile:
        if not PROFILE_PATH.is_file():
            raise FileNotFoundError(f"Cannot reuse missing profile: {PROFILE_PATH}")
        profile = json.loads(PROFILE_PATH.read_text(encoding="utf-8"))
        profiled_variants = profile.get("source", {}).get("variants_profiled", [])
        if not profiled_variants or set(profile.get("variants", {})) != set(profiled_variants):
            raise ValueError("Existing profile is incomplete or has inconsistent variant metadata.")
    elif args.batch_size < 10_000:
        raise ValueError("Use a batch size of at least 10,000 rows.")
    else:
        profile = _profile_all(args.variants, args.batch_size)
    _build_csv_files(profile)
    _build_workbook(
        profile,
        include_hismall_raw=not args.without_hi_small_raw,
    )
    print(f"Profile JSON: {PROFILE_PATH}", flush=True)
    print(
        f"Workbook: {WORKBOOK_PATH if not args.without_hi_small_raw else SUMMARY_WORKBOOK_PATH}",
        flush=True,
    )
    print(f"Transaction template: {TRANSACTION_TEMPLATE}", flush=True)
    print(f"Account template: {ACCOUNT_TEMPLATE}", flush=True)
    print(f"Synthetic development fixture: {SAMPLE_INPUT}", flush=True)
    print(f"Fixture ground truth: {SAMPLE_LABELS}", flush=True)


if __name__ == "__main__":
    main()
